"""Etap 9 — eksport arkusza: formuly, przeliczenie, zgodnosc z silnikiem.

Testy oznaczone `wolne` uruchamiaja LibreOffice. Pomijaj je przez
`pytest -m "not wolne"`, gdy potrzebujesz szybkiego przebiegu.
"""

from decimal import Decimal as D
from pathlib import Path

import openpyxl
import pytest

from sim_kalkulator import arkusz, wrazliwosc
from sim_kalkulator.dane import wczytaj_yaml
from sim_kalkulator.silnik import przelicz

from . import wspolne

DOMYKAJACY = wspolne.KATALOG_PRZYKLADOW / "domykajacy_sie.yaml"

ZAKLADKI = [
    "Zalozenia", "Alokacja", "Pula_spoleczna", "Pula_komunalna",
    "Rekompensata", "Werdykty", "Wrazliwosc", "Podstawy_prawne",
]


@pytest.fixture(scope="module")
def skoroszyt(tmp_path_factory):
    w = wczytaj_yaml(DOMYKAJACY)
    sciezka = tmp_path_factory.mktemp("arkusz") / "wynik.xlsx"
    arkusz.eksportuj(przelicz(w), sciezka, wrazliwosc.build(w, krok=D("0.25")))
    return sciezka


@pytest.fixture(scope="module")
def przeliczony(skoroszyt):
    import importlib.util

    korzen = Path(__file__).resolve().parent.parent
    spec = importlib.util.spec_from_file_location("recalc", korzen / "scripts" / "recalc.py")
    recalc = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(recalc)
    return recalc, recalc.przelicz(skoroszyt)


def komorka(ws, etykieta, kolumna=2):
    for wiersz in ws.iter_rows():
        if wiersz[0].value == etykieta:
            return wiersz[kolumna - 1]
    raise AssertionError(f"Nie ma wiersza '{etykieta}' w zakladce {ws.title}")


class TestStruktura:
    def test_osiem_zakladek_w_kolejnosci_ze_specyfikacji(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        assert wb.sheetnames == ZAKLADKI

    def test_zakladki_projekcji_sa_zbudowane_z_formul(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        for nazwa in ("Alokacja", "Pula_spoleczna", "Pula_komunalna", "Rekompensata", "Werdykty"):
            ws = wb[nazwa]
            formul = sum(
                1
                for wiersz in ws.iter_rows()
                for k in wiersz
                if isinstance(k.value, str) and k.value.startswith("=")
            )
            assert formul > 20, f"{nazwa} ma tylko {formul} formul — arkusz z wartosciami"

    def test_zadna_zakladka_wynikowa_nie_ma_wklejonych_kwot(self, skoroszyt):
        # Poza Zalozeniami, Podstawami_prawnymi i migawka Wrazliwosci kazda liczba
        # ma byc wynikiem formuly, nie wklejona wartoscia.
        wb = openpyxl.load_workbook(skoroszyt)
        for nazwa in ("Alokacja", "Pula_spoleczna", "Pula_komunalna", "Rekompensata", "Werdykty"):
            ws = wb[nazwa]
            wklejone = [
                f"{nazwa}!{k.coordinate}={k.value}"
                for wiersz in ws.iter_rows()
                for k in wiersz
                if isinstance(k.value, (int, float)) and abs(k.value) > 1
            ]
            assert not wklejone, f"Wklejone wartosci: {wklejone[:5]}"

    def test_wejscia_sa_niebieskie_a_odwolania_zielone(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        wejscie = komorka(wb["Zalozenia"], "PUM laczne")
        assert wejscie.font.color.rgb == "FF0000CC"
        odwolanie = komorka(wb["Alokacja"], "Udzial puli komunalnej")
        assert odwolanie.font.color.rgb == "FF006600"

    def test_kluczowe_zalozenia_maja_zolte_wypelnienie(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        for etykieta in ("Udzial puli komunalnej — GLOWNE POKRETLO", "Stopa referencyjna KE (r)"):
            assert komorka(wb["Zalozenia"], etykieta).fill.fgColor.rgb == "FFFFF2A8"

    def test_lata_zapisane_jako_tekst(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        ws = wb["Pula_spoleczna"]
        lata = [
            k for wiersz in ws.iter_rows(min_col=1, max_col=1) for k in wiersz
            if isinstance(k.value, str) and k.value.isdigit()
        ]
        assert lata
        assert all(k.number_format == "@" for k in lata)

    def test_uzyte_sa_tylko_funkcje_ze_standardu_excel_2007(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        zabronione = ("XLOOKUP", "IFS(", "TEXTJOIN", "SWITCH(", "MAXIFS", "MINIFS", "LET(")
        for nazwa in wb.sheetnames:
            for wiersz in wb[nazwa].iter_rows():
                for k in wiersz:
                    if isinstance(k.value, str) and k.value.startswith("="):
                        for funkcja in zabronione:
                            assert funkcja not in k.value.upper(), f"{nazwa}!{k.coordinate}"

    def test_wyszukiwanie_progow_przez_index_match(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        stawka = komorka(wb["Pula_spoleczna"], "Stawka roczna z art. 7c")
        assert "INDEX(" in stawka.value and "MATCH(" in stawka.value


class TestZastrzezenia:
    """Rozdz. 13 — zastrzezenia widoczne w arkuszu, nie tylko w README."""

    def test_zastrzezenia_sa_w_zalozeniach_i_werdyktach(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        for nazwa in ("Zalozenia", "Werdykty"):
            tekst = " ".join(
                str(k.value) for wiersz in wb[nazwa].iter_rows() for k in wiersz if k.value
            )
            assert "Nie zastepuje wyliczenia BGK" in tekst
            assert "Kwestie otwarte" in tekst

    def test_test_3_niesie_zastrzezenie_o_wiazacym_wyliczeniu_banku(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        tekst = " ".join(
            str(k.value) for wiersz in wb["Werdykty"].iter_rows() for k in wiersz if k.value
        )
        assert "WYLICZENIE BGK JEST WIAZACE" in tekst

    def test_kazda_stala_ma_podstawe_w_zakladce_podstaw(self, skoroszyt):
        wb = openpyxl.load_workbook(skoroszyt)
        ws = wb["Podstawy_prawne"]
        wiersze = [
            w for w in ws.iter_rows(min_col=1, max_col=3)
            if isinstance(w[1].value, (int, float)) and w[0].value
        ]
        assert len(wiersze) >= 20
        for w in wiersze:
            assert w[2].value, f"Stala '{w[0].value}' bez podstawy prawnej"


@pytest.mark.wolne
class TestPrzeliczenie:
    def test_zero_bledow_formul(self, przeliczony):
        recalc, plik = przeliczony
        bledy, _ = recalc.zbierz_bledy(plik)
        assert bledy == [], bledy[:10]

    def test_zgodnosc_arkusza_z_silnikiem(self, przeliczony):
        recalc, plik = przeliczony
        assert recalc.porownaj_z_silnikiem(plik, DOMYKAJACY) == []

    def test_werdykty_w_arkuszu_zgodne_z_silnikiem(self, przeliczony):
        recalc, plik = przeliczony
        wb = openpyxl.load_workbook(plik, data_only=True)
        r = przelicz(wczytaj_yaml(DOMYKAJACY))
        ws = wb["Werdykty"]
        assert komorka(ws, "WERDYKT 1").value == r.werdykty.montaz.status
        assert komorka(ws, "WERDYKT 2").value == r.werdykty.zdolnosc_czynszowa.status
        assert komorka(ws, "WERDYKT 3").value == r.werdykty.rekompensata.status
        assert komorka(ws, "CZY MONTAZ SIE DOMYKA").value == (
            "TAK" if r.domyka_sie else "NIE"
        )

    def test_asercja_edb_wychodzi_w_normie(self, przeliczony):
        recalc, plik = przeliczony
        wb = openpyxl.load_workbook(plik, data_only=True)
        assert komorka(wb["Rekompensata"], "Kontrola asercji 0 <= EDB < S", 11).value == "w normie"

    def test_harmonogram_domyka_saldo_kredytu(self, przeliczony):
        recalc, plik = przeliczony
        wb = openpyxl.load_workbook(plik, data_only=True)
        ws = wb["Pula_spoleczna"]
        kwota = komorka(ws, "Kredyt SBC").value
        kapital_razem = komorka(ws, "Razem", 4).value
        assert abs(kapital_razem - kwota) < 0.01


@pytest.mark.wolne
class TestWariantowGruntu:
    """Rozdz. 8 — arkusz ma odtwarzac kanaly gruntowe, nie tylko wariant domyslny.

    Wariant z aportem uruchamia w formulach dwie rzeczy, ktorych nabycie nie
    dotyka: samozwrotny limit z § 12 ust. 7 rozp. 766 i odjecie wkladu rzeczowego
    w tescie kapitalowym.
    """

    @staticmethod
    def _skoroszyt_wariantu(tmp_path, **zmiany):
        import yaml

        from sim_kalkulator.dane import zbuduj

        dane = wspolne.zmien(**zmiany)
        wejscie = tmp_path / "wariant.yaml"
        wejscie.write_text(yaml.safe_dump(dane, allow_unicode=True), encoding="utf-8")
        plik = tmp_path / "wariant.xlsx"
        arkusz.eksportuj(przelicz(zbuduj(dane)), plik)
        return wejscie, plik

    @staticmethod
    def _recalc():
        import importlib.util

        korzen = Path(__file__).resolve().parent.parent
        spec = importlib.util.spec_from_file_location(
            "recalc_wariant", korzen / "scripts" / "recalc.py"
        )
        modul = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(modul)
        return modul

    def test_aport_inwestora_zgadza_sie_z_silnikiem(self, tmp_path):
        recalc = self._recalc()
        wejscie, plik = self._skoroszyt_wariantu(
            tmp_path,
            grunt__pochodzenie="inwestor",
            grunt__forma="aport_inwestora",
            grunt__wartosc=12000000.0,
            grunt__liczba_lokali_dla_gminy=wspolne.USUN,
            grunt__pum_lokali_dla_gminy=wspolne.USUN,
        )
        przeliczony = recalc.przelicz(plik)
        bledy, _ = recalc.zbierz_bledy(przeliczony)
        assert bledy == [], bledy[:10]
        assert recalc.porownaj_z_silnikiem(przeliczony, wejscie) == []

    def test_dzierzawa_zgadza_sie_z_silnikiem(self, tmp_path):
        recalc = self._recalc()
        wejscie, plik = self._skoroszyt_wariantu(
            tmp_path,
            grunt__forma="dzierzawa",
            grunt__liczba_lokali_dla_gminy=wspolne.USUN,
            grunt__pum_lokali_dla_gminy=wspolne.USUN,
        )
        przeliczony = recalc.przelicz(plik)
        bledy, _ = recalc.zbierz_bledy(przeliczony)
        assert bledy == [], bledy[:10]
        assert recalc.porownaj_z_silnikiem(przeliczony, wejscie) == []

    def test_uzytkowanie_wieczyste_zgadza_sie_z_silnikiem(self, tmp_path):
        # Po usunieciu aportu gminy (pakiet nr 2, rozdz. 11) to jedyna forma,
        # przy ktorej grunt jest przychodem uslugi publicznej — czyli jedyna,
        # ktora uruchamia w arkuszu galaz kanalu C.
        recalc = self._recalc()
        wejscie, plik = self._skoroszyt_wariantu(
            tmp_path,
            grunt__forma="uzytkowanie_wieczyste",
            grunt__liczba_lokali_dla_gminy=wspolne.USUN,
            grunt__pum_lokali_dla_gminy=wspolne.USUN,
        )
        przeliczony = recalc.przelicz(plik)
        bledy, _ = recalc.zbierz_bledy(przeliczony)
        assert bledy == [], bledy[:10]
        assert recalc.porownaj_z_silnikiem(przeliczony, wejscie) == []


@pytest.mark.wolne
class TestRecznaKontrolaFormul:
    """Rozdz. 8.4 — wyrywkowa kontrola formul w zakladkach projekcji."""

    @staticmethod
    @pytest.fixture(scope="class")
    def dane(przeliczony):
        _, plik = przeliczony
        return openpyxl.load_workbook(plik, data_only=True)

    def test_przychod_netto_roku_1_liczony_recznie(self, dane):
        ws = dane["Pula_spoleczna"]
        czynsz = komorka(ws, "Czynsz zakladany").value
        pum = dane["Alokacja"]["C" + str(_wiersz(dane["Alokacja"], "PUM"))].value
        naglowek = _wiersz_naglowka(ws, "Rok", "Indeks czynszu")
        rok1 = naglowek + 1
        recznie = czynsz * 12 * pum
        assert abs(ws[f"E{rok1}"].value - recznie) < 0.01
        assert abs(ws[f"F{rok1}"].value - recznie * 0.05) < 0.01
        assert abs(ws[f"G{rok1}"].value - recznie * 0.95) < 0.01

    def test_indeksacja_roku_5_liczona_recznie(self, dane):
        ws = dane["Pula_spoleczna"]
        naglowek = _wiersz_naglowka(ws, "Rok", "Indeks czynszu")
        rok5 = naglowek + 5
        assert abs(ws[f"B{rok5}"].value - 1.030 ** 4) < 1e-9
        assert abs(ws[f"C{rok5}"].value - 1.035 ** 4) < 1e-9

    def test_dscr_roku_3_liczony_recznie(self, dane):
        ws = dane["Pula_spoleczna"]
        naglowek = _wiersz_naglowka(ws, "Rok", "Indeks czynszu")
        rok3 = naglowek + 3
        # H..M: eksploatacja, odpis, ubezpieczenie, zarzad, oplata za grunt, rata.
        pokrycie = sum(ws[f"{k}{rok3}"].value for k in "HIJKLM")
        assert abs(ws[f"O{rok3}"].value - pokrycie) < 0.01
        assert abs(ws[f"P{rok3}"].value - ws[f"G{rok3}"].value / pokrycie) < 1e-9

    def test_dyskonto_w_rekompensacie_liczone_recznie(self, dane):
        ws = dane["Rekompensata"]
        naglowek = _wiersz_naglowka(ws, "Rok", "Koszty biezace")
        rok4 = naglowek + 4
        rb = dane["Zalozenia"]["B" + str(_wiersz(dane["Zalozenia"], "Stopa bazowa KE (rb)"))].value
        assert abs(ws[f"J{rok4}"].value - (1 + rb) ** 3) < 1e-9
        assert abs(
            ws[f"K{rok4}"].value
            - (ws[f"F{rok4}"].value - ws[f"I{rok4}"].value) / ws[f"J{rok4}"].value
        ) < 0.01

    def test_pula_komunalna_nie_ma_obslugi_dlugu_ani_rezerwy(self, dane):
        ws = dane["Pula_komunalna"]
        naglowek = _wiersz_naglowka(ws, "Rok", "Indeks czynszu")
        for rok in range(1, 6):
            assert ws[f"M{naglowek + rok}"].value == 0
            assert ws[f"N{naglowek + rok}"].value == 0


@pytest.mark.wolne
class TestZmianaZalozenia:
    """Rozdz. 8.1 — zmiana zalozenia w zakladce wejsciowej przelicza caly skoroszyt."""

    def test_zmiana_pokretla_przelicza_wszystkie_zakladki(self, skoroszyt, tmp_path):
        import importlib.util

        korzen = Path(__file__).resolve().parent.parent
        spec = importlib.util.spec_from_file_location("recalc", korzen / "scripts" / "recalc.py")
        recalc = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(recalc)

        przed = openpyxl.load_workbook(recalc.przelicz(skoroszyt), data_only=True)
        wartosci_przed = {
            "koszty_spol": komorka(przed["Alokacja"], "Koszty przedsiewziecia", 3).value,
            "grant_kom": komorka(przed["Alokacja"], "GRANT", 4).value,
            "wklad": komorka(przed["Werdykty"], "WYMAGANY WKLAD WLASNY").value,
            "kn_kom": komorka(przed["Rekompensata"], "KOSZTY NETTO (KN)", 11).value,
        }

        # Zmieniamy WYLACZNIE jedna komorke wejsciowa w zakladce Zalozenia.
        wb = openpyxl.load_workbook(skoroszyt)
        cel = komorka(wb["Zalozenia"], "Udzial puli komunalnej — GLOWNE POKRETLO")
        assert cel.value == 0.30
        cel.value = 0.55
        zmieniony = tmp_path / "zmieniony.xlsx"
        wb.save(zmieniony)

        po = openpyxl.load_workbook(recalc.przelicz(zmieniony), data_only=True)
        wartosci_po = {
            "koszty_spol": komorka(po["Alokacja"], "Koszty przedsiewziecia", 3).value,
            "grant_kom": komorka(po["Alokacja"], "GRANT", 4).value,
            "wklad": komorka(po["Werdykty"], "WYMAGANY WKLAD WLASNY").value,
            "kn_kom": komorka(po["Rekompensata"], "KOSZTY NETTO (KN)", 11).value,
        }
        for klucz in wartosci_przed:
            assert wartosci_przed[klucz] != wartosci_po[klucz], (
                f"'{klucz}' nie przeliczyl sie po zmianie zalozenia — arkusz jest martwy"
            )

    def test_zmieniony_arkusz_zgadza_sie_z_silnikiem(self, skoroszyt, tmp_path):
        import importlib.util

        korzen = Path(__file__).resolve().parent.parent
        spec = importlib.util.spec_from_file_location("recalc", korzen / "scripts" / "recalc.py")
        recalc = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(recalc)

        wb = openpyxl.load_workbook(skoroszyt)
        komorka(wb["Zalozenia"], "Udzial puli komunalnej — GLOWNE POKRETLO").value = 0.55
        zmieniony = tmp_path / "zmieniony.xlsx"
        wb.save(zmieniony)

        po = openpyxl.load_workbook(recalc.przelicz(zmieniony), data_only=True)
        r = przelicz(wczytaj_yaml(DOMYKAJACY).z_udzialem_komunalnym(D("0.55")))
        assert abs(
            komorka(po["Werdykty"], "WYMAGANY WKLAD WLASNY").value
            - float(r.finansowanie.wklad_wlasny_wymagany)
        ) < 0.01
        assert abs(
            komorka(po["Alokacja"], "GRANT", 3).value - float(r.granty.spoleczna.kwota)
        ) < 0.01


@pytest.fixture(scope="module")
def skoroszyt_miedzy_progami(tmp_path_factory):
    from tests.test_audyt import MIEDZY_PROGAMI

    w = wspolne.wejscie(**MIEDZY_PROGAMI)
    sciezka = tmp_path_factory.mktemp("prog") / "wynik.xlsx"
    arkusz.eksportuj(przelicz(w), sciezka, wrazliwosc.build(w, krok=D("0.5")))
    return sciezka


class TestProguTolerancjiWArkuszu:
    """Pakiet nr 2, rozdz. 2+3 — prog jest komorka, a werdykt idzie za profilem."""

    def _recalc(self):
        import importlib.util

        korzen = Path(__file__).resolve().parent.parent
        spec = importlib.util.spec_from_file_location(
            "recalc", korzen / "scripts" / "recalc.py"
        )
        recalc = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(recalc)
        return recalc

    def test_arkusz_zgadza_sie_z_silnikiem_przy_progu_nizszym(
        self, skoroszyt_miedzy_progami
    ):
        from tests.test_audyt import MIEDZY_PROGAMI

        recalc = self._recalc()
        po = openpyxl.load_workbook(
            recalc.przelicz(skoroszyt_miedzy_progami), data_only=True
        )
        r = przelicz(wspolne.wejscie(**MIEDZY_PROGAMI))
        assert komorka(po["Werdykty"], "WERDYKT 3").value == "nie przechodzi"
        assert r.werdykty.rekompensata.przechodzi is False
        najgorsza = komorka(
            po["Rekompensata"], "Najgorszy rok — nadwyzka wzgledna", 11
        ).value
        assert abs(
            najgorsza - float(r.rekompensata.spoleczna.nadwyzka_wzgledna_najgorsza)
        ) < 0.0001

    def test_podmiana_progu_w_arkuszu_odwraca_werdykt(
        self, skoroszyt_miedzy_progami, tmp_path
    ):
        # Prog jest ZALOZENIEM, wiec ma byc komorka do podmiany, a nie stala
        # wpisana w formule. Test sprawdza obie rzeczy naraz: ze da sie ja
        # podmienic i ze arkusz faktycznie przelicza od niej werdykt.
        from tests.test_audyt import MIEDZY_PROGAMI

        recalc = self._recalc()
        wb = openpyxl.load_workbook(skoroszyt_miedzy_progami)
        cel = komorka(wb["Rekompensata"], "Prog tolerancji nadwyzki — udzial", 11)
        assert cel.value == 0.10
        cel.value = 0.20
        zmieniony = tmp_path / "prog_wyzszy.xlsx"
        wb.save(zmieniony)

        po = openpyxl.load_workbook(recalc.przelicz(zmieniony), data_only=True)
        assert komorka(po["Werdykty"], "WERDYKT 3").value == "przechodzi"
        # ...i silnik przy tym samym zalozeniu mowi to samo.
        r = przelicz(
            wspolne.wejscie(
                **MIEDZY_PROGAMI,
                przelaczniki__prog_tolerancji_przy_dwoch_instrumentach="wyzszy",
            )
        )
        assert r.werdykty.rekompensata.przechodzi is True


def _wiersz(ws, etykieta):
    for wiersz in ws.iter_rows(min_col=1, max_col=1):
        if wiersz[0].value == etykieta:
            return wiersz[0].row
    raise AssertionError(f"Nie ma wiersza '{etykieta}' w {ws.title}")


def _wiersz_naglowka(ws, pierwsza, druga):
    for wiersz in ws.iter_rows(min_col=1, max_col=2):
        if wiersz[0].value == pierwsza and wiersz[1].value == druga:
            return wiersz[0].row
    raise AssertionError(f"Nie ma naglowka '{pierwsza}/{druga}' w {ws.title}")
