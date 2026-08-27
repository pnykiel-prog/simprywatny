"""Eksport skoroszytu XLSX — formuly, nie wartosci.

Arkusz z wklejonymi liczbami jest nieweryfikowalny i nie nadaje sie do
dokumentacji wniosku. Kazda komorka wynikowa jest formula odwolujaca sie do
komorek zalozen: zmiana zalozenia w zakladce `Zalozenia` przelicza caly
skoroszyt.

Konwencja kolorow (rozdz. 8.3 specyfikacji):
  * wejscia i dzwignie scenariuszowe — tekst niebieski,
  * formuly — czarny,
  * odwolania miedzyzakladkowe — zielony,
  * kluczowe zalozenia do uzupelnienia — zolte wypelnienie.

Funkcje ograniczone do standardu Excel 2007. Do wyszukiwan INDEX/MATCH, nie XLOOKUP.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import prawo
from . import zalozenia as _katalog_zalozen
from .dane import (
    FormaGruntu,
    MetodaRozsadnegoZysku,
    TrybKredytu,
    UjecieKosztowInwestycyjnych,
)
from .silnik import Wynik
from .wrazliwosc import Analiza
from .waluta import ZERO

# --- formaty liczbowe ------------------------------------------------------
# Separator grup renderuje sie wg ustawien jezykowych; w polskiej lokalizacji
# daje to zapis "1 234 zl" wymagany w rozdz. 8.3.
KWOTA = '#,##0" zl"'
KWOTA_GROSZE = '#,##0.00" zl"'
STAWKA = '#,##0.00" zl/m2/mies."'
PROCENT = "0.0%"
PROCENT_DOKLADNY = "0.00%"
WSKAZNIK = "0.000"
TEKST = "@"
POWIERZCHNIA = '#,##0.0" m2"'

# --- kolory ----------------------------------------------------------------
NIEBIESKI = Font(color="FF0000CC")
CZARNY = Font(color="FF000000")
ZIELONY = Font(color="FF006600")
NAGLOWEK = Font(bold=True, color="FF000000", size=12)
SEKCJA = Font(bold=True, color="FF333333")
ZOLTE = PatternFill("solid", fgColor="FFFFF2A8")
SZARE = PatternFill("solid", fgColor="FFEFEFEF")
CZERWONY = Font(bold=True, color="FFAA0000")
ZIELONY_BOLD = Font(bold=True, color="FF006600")

RAMKA_DOL = Border(bottom=Side(style="thin", color="FFAAAAAA"))


class Rejestr:
    """Adresy komorek do budowania formul miedzyzakladkowych."""

    def __init__(self) -> None:
        self._adresy: Dict[str, str] = {}
        self._wiersze: Dict[str, int] = {}

    def zapisz(self, klucz: str, arkusz: str, komorka: str) -> str:
        adres = f"'{arkusz}'!{komorka}"
        self._adresy[klucz] = adres
        return adres

    def __getitem__(self, klucz: str) -> str:
        if klucz not in self._adresy:
            raise KeyError(
                f"Brak adresu '{klucz}' w rejestrze arkusza. "
                "Formula odwolywalaby sie w prozne miejsce."
            )
        return self._adresy[klucz]

    def ma(self, klucz: str) -> bool:
        return klucz in self._adresy

    def zapisz_wiersz(self, klucz: str, numer: int) -> int:
        """Numer wiersza bez prefiksu arkusza — do sklejania adresow typu H62."""
        self._wiersze[klucz] = numer
        return numer

    def wiersz(self, klucz: str) -> int:
        if klucz not in self._wiersze:
            raise KeyError(f"Brak numeru wiersza '{klucz}' w rejestrze arkusza.")
        return self._wiersze[klucz]

    def ma_wiersz(self, klucz: str) -> bool:
        return klucz in self._wiersze


def _bezwzgledny(kolumna: str, wiersz: int) -> str:
    return f"${kolumna}${wiersz}"


def _naglowek(ws: Worksheet, wiersz: int, tytul: str, podtytul: str = "") -> int:
    ws.cell(row=wiersz, column=1, value=tytul).font = NAGLOWEK
    if podtytul:
        ws.cell(row=wiersz + 1, column=1, value=podtytul).font = Font(
            italic=True, color="FF666666", size=9
        )
        return wiersz + 3
    return wiersz + 2


def _sekcja(ws: Worksheet, wiersz: int, tytul: str) -> int:
    komorka = ws.cell(row=wiersz, column=1, value=tytul)
    komorka.font = SEKCJA
    komorka.fill = SZARE
    for kol in range(2, 7):
        ws.cell(row=wiersz, column=kol).fill = SZARE
    return wiersz + 1


def _szerokosci(ws: Worksheet, szerokosci: Sequence[Tuple[str, int]]) -> None:
    for kolumna, szerokosc in szerokosci:
        ws.column_dimensions[kolumna].width = szerokosc


# ===========================================================================
# Zakladka Zalozenia — kazdy parametr w osobnej komorce, ze zrodlem i data
# ===========================================================================

def _zalozenia(wb: Workbook, wynik: Wynik, rej: Rejestr) -> None:
    w = wynik.wejscie
    u = wynik.grunt
    ws = wb.create_sheet("Zalozenia")
    _szerokosci(ws, [("A", 46), ("B", 20), ("C", 18), ("D", 58), ("E", 52)])

    wiersz = _naglowek(
        ws,
        1,
        f"Zalozenia — {w.projekt.nazwa}",
        "Tekst niebieski = wejscie do edycji. Zmiana dowolnej komorki w kolumnie B "
        "przelicza caly skoroszyt. Zolte wypelnienie = parametr zewnetrzny do potwierdzenia przed naborem.",
    )
    for kol, tytul in enumerate(
        ["Parametr", "Wartosc", "Jednostka", "Zrodlo / uwaga", "Podstawa prawna"], start=1
    ):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True)
        komorka.border = RAMKA_DOL
    wiersz += 1

    def wejscie(
        klucz: str,
        nazwa: str,
        wartosc,
        jednostka: str = "",
        zrodlo: str = "",
        podstawa: str = "",
        format_liczby: str = KWOTA_GROSZE,
        zolte: bool = False,
    ) -> None:
        nonlocal wiersz
        ws.cell(row=wiersz, column=1, value=nazwa)
        komorka = ws.cell(row=wiersz, column=2)
        komorka.value = float(wartosc) if isinstance(wartosc, Decimal) else wartosc
        komorka.font = NIEBIESKI
        komorka.number_format = format_liczby
        if zolte:
            komorka.fill = ZOLTE
        ws.cell(row=wiersz, column=3, value=jednostka)
        ws.cell(row=wiersz, column=4, value=zrodlo)
        ws.cell(row=wiersz, column=5, value=podstawa).font = Font(size=9, color="FF666666")
        rej.zapisz(klucz, "Zalozenia", _bezwzgledny("B", wiersz))
        wiersz += 1

    # --- projekt ---
    wiersz = _sekcja(ws, wiersz, "PROJEKT")
    for klucz, nazwa, wartosc in (
        ("projekt.nazwa", "Nazwa", w.projekt.nazwa),
        ("projekt.gmina", "Gmina", w.projekt.gmina),
        ("projekt.wojewodztwo", "Wojewodztwo (wskaznik wartosci odtworzeniowej)", w.projekt.wojewodztwo),
    ):
        wejscie(klucz, nazwa, wartosc, format_liczby=TEKST)

    # --- powierzchnie ---
    wiersz = _sekcja(ws, wiersz, "POWIERZCHNIE")
    p = w.powierzchnie
    wejscie("pum_laczne", "PUM laczne", p.pum_laczne, "m2", format_liczby=POWIERZCHNIA)
    wejscie("liczba_lokali", "Liczba lokali", p.liczba_lokali, "szt.", format_liczby="0")
    wejscie(
        "liczba_kondygnacji", "Liczba kondygnacji naziemnych", p.liczba_kondygnacji, "szt.",
        zrodlo="Od 3 kondygnacji dzwigi osobowe sa obowiazkowe.",
        podstawa="rozp. MIiR z 4.03.2019, Dz.U. 2019 poz. 457", format_liczby="0",
    )
    wejscie(
        "udzial_komunalny", "Udzial puli komunalnej — GLOWNE POKRETLO",
        p.udzial_puli_komunalnej, "udzial",
        zrodlo="Dzwignia scenariuszowa. Tresc negocjacji z gmina.",
        format_liczby=PROCENT, zolte=True,
    )

    # --- koszty ---
    wiersz = _sekcja(ws, wiersz, "KOSZTY PRZEDSIEWZIECIA")
    k = w.koszty
    wejscie("koszt_budowy_m2", "Koszt budowy na m2 PUM", k.koszt_budowy_na_m2, "zl/m2")
    wejscie("infrastruktura", "Infrastruktura", k.infrastruktura, "zl")
    wejscie("projekt_i_nadzor", "Projekt i nadzor", k.projekt_i_nadzor, "zl")
    wejscie("koszty_ogolne", "Koszty ogolne", k.koszty_ogolne, "zl")
    wejscie("rezerwa", "Rezerwa", k.rezerwa, "zl")
    wejscie("dzwigi", "Dzwigi osobowe", k.dzwigi, "zl",
            podstawa="rozp. MIiR z 4.03.2019, Dz.U. 2019 poz. 457")
    wejscie(
        "vat_odliczalny", "VAT odliczalny", k.vat_odliczalny, "TRUE/FALSE",
        zrodlo="FALSE podnosi podstawe grantu i kwote do sfinansowania o VAT.",
        podstawa="art. 13 ust. 3 ustawy z 8.12.2006", format_liczby=TEKST,
    )
    wejscie("stawka_vat", "Stawka VAT", k.stawka_vat, "udzial",
            zrodlo="Stosowana tylko gdy VAT nieodliczalny.", format_liczby=PROCENT)

    # --- grunt ---
    wiersz = _sekcja(ws, wiersz, "GRUNT")
    wejscie("grunt_wartosc", "Wartosc gruntu", w.grunt.wartosc, "zl",
            zrodlo="Z operatu szacunkowego.",
            podstawa="art. 13 ust. 1 pkt 1 ustawy z 8.12.2006", zolte=True)
    wejscie("grunt_pochodzenie", "Pochodzenie dzialki", w.grunt.pochodzenie.value, "",
            zrodlo="Poziom 1 wyboru: inwestor / rynek_prywatny / gmina.",
            format_liczby=TEKST, zolte=True)
    wejscie("grunt_forma", "Forma wniesienia", w.grunt.forma.value, "",
            zrodlo="Poziom 2. Rozstrzyga cztery kanaly naraz — patrz matryca skutkow.",
            podstawa="art. 5 ust. 7 pkt 7 / art. 5 ust. 9 pkt 4 / § 12 ust. 7 rozp. 766",
            format_liczby=TEKST, zolte=True)
    wejscie("grunt_pasmo_45", "Kanal A: forma daje pasmo ponad prog", u.pasmo_45, "TRUE/FALSE",
            zrodlo="Wyliczone z formy. FALSE scina wsparcie do progu gruntowego.",
            podstawa="art. 13 ust. 1 pkt 1 ustawy z 8.12.2006", format_liczby=TEKST)
    wejscie("grunt_w_kosztach", "Kanal B: wartosc wchodzi do kosztow",
            u.wartosc_w_kosztach > 0, "TRUE/FALSE",
            zrodlo="Wyliczone z formy. Przy dzierzawie FALSE — liczy sie oplata roczna.",
            podstawa="art. 5 ust. 7 pkt 7 i ust. 8 ustawy z 8.12.2006", format_liczby=TEKST)
    wejscie("grunt_aport", "Kanal B: wklad niepieniezny (limit 20%)",
            u.limit_aportowy_dotyczy, "TRUE/FALSE",
            zrodlo="Wyliczone z formy. Limit dziala tylko w sciezce kredytowej.",
            podstawa="§ 12 ust. 7 rozp. t.j. Dz.U. 2021 poz. 766", format_liczby=TEKST)
    wejscie("grunt_przychod_uoig", "Kanal C: wartosc jest przychodem UOIG",
            u.przychod_uoig, "TRUE/FALSE",
            zrodlo="Wyliczone z formy i przelacznikow 9.1-9.2. TRUE obniza dopuszczalna pomoc.",
            podstawa="art. 5 ust. 9 pkt 4 ustawy z 8.12.2006", format_liczby=TEKST)
    wejscie("grunt_gotowka", "Kanal D: wartosc gruntu placona gotowka",
            u.wydatek_gotowkowy > 0, "TRUE/FALSE",
            zrodlo="Wyliczone z formy. FALSE oznacza wklad rzeczowy — koszt bez wydatku.",
            format_liczby=TEKST)
    wejscie("grunt_do_pasma", "Kanal A: wartosc prawa do limitu", u.wartosc_do_pasma, "zl",
            zrodlo="Wartosc z operatu albo cena po bonifikacie — kwestia 9.3.",
            podstawa="art. 13 ust. 1 pkt 1 ustawy z 8.12.2006")
    wejscie("grunt_oplata_roczna", "Kanal D: oplata roczna za grunt", u.oplata_roczna, "zl/rok",
            zrodlo="Dzierzawa albo uzytkowanie wieczyste. Koszt biezacy, nie kapitalowy.",
            zolte=bool(u.oplata_roczna))
    wejscie("grunt_pum_gminy", "Lokale dla gminy — PUM", u.pum_dla_gminy, "m2",
            zrodlo="Tryb 'lokal za grunt'. Ta powierzchnia nie przynosi czynszu.",
            podstawa="ustawa z 16.12.2020, Dz.U. 2021 poz. 223", zolte=bool(u.pum_dla_gminy))

    # --- pula spoleczna ---
    wiersz = _sekcja(ws, wiersz, "PULA SPOLECZNA")
    ps = w.pula_spoleczna
    wejscie("kredyt_udzial", "Kredyt SBC — udzial docelowy", ps.kredyt.udzial_docelowy, "udzial",
            zrodlo="Maksimum 80% kosztow przedsiewziecia.",
            podstawa="art. 15b ust. 2 ustawy z 26.10.1995", format_liczby=PROCENT)
    wejscie("kredyt_rp", "Kredyt SBC — oprocentowanie (rp)", ps.kredyt.oprocentowanie, "udzial",
            format_liczby=PROCENT_DOKLADNY)
    wejscie("kredyt_n", "Kredyt SBC — okres w latach (N)", ps.kredyt.okres_lat, "lata",
            zrodlo="Maksimum 30 lat wliczajac karencje.",
            podstawa="art. 15b ust. 3 ustawy z 26.10.1995", format_liczby="0")
    wejscie("kredyt_t", "Kredyt SBC — karencja w latach (T)", ps.kredyt.karencja_lat, "lata",
            format_liczby="0")
    wejscie("partycypacja_stawka", "Partycypacja — stawka",
            ps.partycypacja.stawka_procent_kosztu_lokalu, "udzial",
            zrodlo="Maksimum 30% dla osoby fizycznej przy finansowaniu zwrotnym.",
            podstawa="art. 29a ust. 2 ustawy z 26.10.1995", format_liczby=PROCENT)
    wejscie("partycypacja_rotacja", "Partycypacja — rotacja roczna",
            ps.partycypacja.rotacja_roczna, "udzial",
            zrodlo="Do rezerwy na zwrot partycypacji.",
            podstawa="art. 29a ust. 3 ustawy z 26.10.1995", format_liczby=PROCENT)
    wejscie("czynsz_spoleczny", "Czynsz zakladany", ps.czynsz_zakladany_m2_mies,
            "zl/m2/mies.", format_liczby=STAWKA, zolte=True)
    wejscie("bonus_spoleczny", "Bonus rewitalizacyjny / Za zyciem",
            ps.bonus_rewitalizacyjny, "TRUE/FALSE",
            zrodlo="Niedostepny przy finansowaniu zwrotnym.",
            podstawa="art. 13 ust. 4 ustawy z 8.12.2006", format_liczby=TEKST)

    # --- pula komunalna ---
    wiersz = _sekcja(ws, wiersz, "PULA KOMUNALNA")
    wejscie("czynsz_komunalny", "Czynsz placony przez gmine",
            w.pula_komunalna.czynsz_placony_przez_gmine_m2_mies, "zl/m2/mies.",
            format_liczby=STAWKA, zolte=True)
    wejscie("bonus_komunalny", "Bonus rewitalizacyjny / Za zyciem",
            w.pula_komunalna.bonus_rewitalizacyjny, "TRUE/FALSE",
            podstawa="art. 13 ust. 4 ustawy z 8.12.2006", format_liczby=TEKST)
    ws.cell(row=wiersz, column=1,
            value="Kredyt i partycypacja sa w tej puli niedopuszczalne.").font = Font(
        italic=True, size=9, color="FF666666")
    ws.cell(row=wiersz, column=5,
            value="art. 5a ust. 3 ustawy z 8.12.2006").font = Font(size=9, color="FF666666")
    wiersz += 1

    # --- eksploatacja ---
    wiersz = _sekcja(ws, wiersz, "EKSPLOATACJA")
    e = w.eksploatacja
    wejscie("koszt_ekspl_m2", "Koszt eksploatacji", e.koszt_eksploatacji_m2_rok, "zl/m2/rok")
    wejscie("odpis_remontowy_m2", "Odpis remontowy", e.odpis_remontowy_m2_rok, "zl/m2/rok")
    wejscie("ubezpieczenie", "Ubezpieczenie", e.ubezpieczenie_rocznie, "zl/rok")
    wejscie("zarzad", "Koszty stale zarzadu", e.koszty_stale_zarzadu_rocznie, "zl/rok",
            podstawa="w sciezce kredytowej: § 12 ust. 3 rozp. Dz.U. 2021 poz. 766")
    wejscie("pustostany", "Pustostany", e.pustostany_procent, "udzial", format_liczby=PROCENT)
    wejscie("indeks_kosztow", "Indeksacja kosztow", e.indeksacja_kosztow_rocznie, "rocznie",
            format_liczby=PROCENT_DOKLADNY)
    wejscie("indeks_czynszu", "Indeksacja czynszu", e.indeksacja_czynszu_rocznie, "rocznie",
            format_liczby=PROCENT_DOKLADNY)

    # --- parametry zewnetrzne ---
    wiersz = _sekcja(ws, wiersz, "PARAMETRY ZEWNETRZNE — sprawdzic przed kazdym naborem")
    pz = w.parametry_zewnetrzne
    for klucz, nazwa, wartosc, jednostka, fmt in (
        ("wartosc_odtworzeniowa", "Wartosc odtworzeniowa 1 m2", pz.wartosc_odtworzeniowa_m2, "zl/m2", KWOTA_GROSZE),
        ("rb", "Stopa bazowa KE (rb)", pz.stopa_bazowa_ke, "udzial", PROCENT_DOKLADNY),
        ("r", "Stopa referencyjna KE (r)", pz.stopa_referencyjna_ke, "udzial", PROCENT_DOKLADNY),
        ("rd", "Stopa dyskontowa (rd)", pz.stopa_dyskontowa, "udzial", PROCENT_DOKLADNY),
        ("irs", "Stopa IRS BGK", pz.stopa_irs_bgk, "udzial", PROCENT_DOKLADNY),
        ("waloryzacja_part", "Waloryzacja partycypacji (wskaznik GUS)",
         pz.waloryzacja_partycypacji_rocznie, "rocznie", PROCENT_DOKLADNY),
        ("okres_amortyzacji", "Okres amortyzacji budynkow",
         pz.okres_amortyzacji_budynkow_lat, "lata", "0"),
    ):
        wejscie(klucz, nazwa, wartosc, jednostka,
                zrodlo=pz.zrodla.get(klucz.replace("wartosc_odtworzeniowa", "wartosc_odtworzeniowa_m2"), "")
                or _zrodlo_parametru(pz, nazwa),
                format_liczby=fmt, zolte=True)
    wejscie(
        "bufor_dscr", "Minimalny wskaznik pokrycia obslugi dlugu",
        pz.minimalny_wskaznik_pokrycia_obslugi_dlugu, "krotnosc",
        zrodlo=(
            "ZALOZENIE, nie odczyt przepisu — ani rozp. o finansowaniu zwrotnym, "
            "ani informator BGK nie podaja wymaganego pokrycia. Do potwierdzenia w BGK. "
            "Wymiaruje kredyt maksymalny; prog testu 2 pozostaje 1,00."
        ),
        format_liczby="0.00", zolte=True,
    )
    wejscie("data_parametrow", "Data parametrow", pz.data_parametrow.isoformat(), "",
            zrodlo=f"Prog starzenia: {prawo.PARAMETRY_MAKSYMALNY_WIEK_MIESIECY} miesiecy.",
            format_liczby=TEKST, zolte=True)

    # --- rekompensata i inwestor ---
    wiersz = _sekcja(ws, wiersz, "REKOMPENSATA — SKLADNIKI DODATKOWE")
    wejscie("rfrm", "Wsparcie z RFRM", w.rekompensata.wsparcie_rfrm, "zl",
            podstawa="§ 7 ust. 7 rozp. Dz.U. 2025 poz. 1897")
    wejscie("dokumentacja_bgk", "Wartosc prawa do dokumentacji BGK",
            w.rekompensata.wartosc_dokumentacji_bgk, "zl",
            zrodlo="Dokumentacja z zasobu Banku nie jest darmowa w sensie limitu.",
            podstawa="§ 7 ust. 7 rozp. Dz.U. 2025 poz. 1897")
    wejscie("rz_kwota", "Rozsadny zysk — kwota wprost",
            w.rekompensata.rozsadny_zysk_kwota if w.rekompensata.rozsadny_zysk_kwota is not None else "",
            "zl", zrodlo="Uzywane tylko przy metodzie 'kwota_wprost'.")

    wiersz = _sekcja(ws, wiersz, "INWESTOR")
    wejscie("wklad_dostepny", "Zadeklarowany kapital inwestora (opcjonalnie)",
            w.inwestor.dostepny_wklad_wlasny if w.inwestor.zadeklarowany else "", "zl",
            zrodlo="Punkt odniesienia. Nie jest potrzebny do obliczenia.", zolte=True)

    # --- rynek najmu ---
    wiersz = _sekcja(ws, wiersz, "RYNEK NAJMU — SUFIT FAKTYCZNY, NIE WYLICZANY PRZEZ SILNIK")
    ry = w.rynek
    wejscie(
        "czynsz_rynkowy", "Czynsz rynkowy w tej miejscowosci",
        ry.czynsz_rynkowy_m2_mies if ry.podano else "",
        "zl/m2/mies.",
        zrodlo=(
            "Poziom akceptowany przez rynek najmu. Silnik go NIE WYLICZA — to obserwacja. "
            "Pusta komorka znaczy 'nie podano', nie zero. Dotyczy wylacznie puli spolecznej: "
            "w komunalnej najemca jest gmina."
        ),
        format_liczby=STAWKA, zolte=True,
    )
    wejscie(
        "rynek_zrodlo", "Skad ta stawka", ry.zrodlo or "", "",
        zrodlo="Obowiazkowe, gdy podano stawke. Bez tego wyniku nie da sie odtworzyc.",
        format_liczby=TEKST, zolte=ry.podano,
    )
    wejscie(
        "rynek_data", "Data obserwacji", ry.data.isoformat() if ry.data else "", "",
        zrodlo="Rynek najmu zmienia sie szybciej niz wskazniki ustawowe.",
        format_liczby=TEKST,
    )

    # --- przelaczniki ---
    wiersz = _sekcja(ws, wiersz, "PRZELACZNIKI — KWESTIE OTWARTE, WARTOSCI DOMYSLNE SA ZALOZENIAMI")
    pzz = w.przelaczniki
    wejscie("sw_hybryda", "Hybryda jako jedno przedsiewziecie",
            pzz.hybryda_jako_jedno_przedsiewziecie, "TRUE/FALSE",
            zrodlo="Kwestia otwarta 10.1. Domyslnie dwa odrebne przedsiewziecia.",
            podstawa="art. 13 ust. 1a w zw. z art. 5a ust. 1 i 3", format_liczby=TEKST, zolte=True)
    wejscie("sw_lokal_za_grunt_przychod", "9.1 Lokal za grunt jest przychodem UOIG",
            pzz.lokal_za_grunt_jest_przychodem_uoig, "TRUE/FALSE",
            zrodlo="Domyslnie FALSE — to nabycie, a nie wniesienie przez JST. "
                   "ZALOZENIE do potwierdzenia w BGK.",
            podstawa="art. 5 ust. 9 pkt 4 ustawy z 8.12.2006", format_liczby=TEKST, zolte=True)
    wejscie("sw_uw_przychod", "9.2 Uzytkowanie wieczyste jest przychodem UOIG",
            pzz.uzytkowanie_wieczyste_jest_przychodem_uoig, "TRUE/FALSE",
            zrodlo="Domyslnie TRUE — wariant ostrozniejszy. ZALOZENIE do potwierdzenia w BGK.",
            podstawa="art. 5 ust. 9 pkt 4 ustawy z 8.12.2006", format_liczby=TEKST, zolte=True)
    wejscie("sw_pasmo_operat", "9.3 Pasmo liczone od wartosci z operatu",
            pzz.pasmo_liczone_od_wartosci_z_operatu, "TRUE/FALSE",
            zrodlo="Domyslnie TRUE — przepis mowi o wartosci prawa, nie o cenie nabycia. "
                   "ZALOZENIE do potwierdzenia w BGK.",
            podstawa="art. 13 ust. 1 pkt 1 ustawy z 8.12.2006", format_liczby=TEKST, zolte=True)
    wejscie("sw_remont", "Remont i przebudowa zamiast budowy",
            pzz.remont_i_przebudowa, "TRUE/FALSE",
            zrodlo="Podnosi limit czynszu z art. 7c do 5,0%.",
            podstawa="art. 5 ust. 1 pkt 2 ustawy z 8.12.2006", format_liczby=TEKST)
    wejscie("sw_pustostany_kom", "Pustostany takze w puli komunalnej",
            pzz.pustostany_takze_w_puli_komunalnej, "TRUE/FALSE",
            zrodlo="Zalozenie modelowe. Sprawdz projekt umowy z gmina.",
            format_liczby=TEKST, zolte=True)
    wejscie("sw_ujecie_inwest", "Ujecie nakladu inwestycyjnego w KN",
            pzz.koszty_inwestycyjne_w_kn.value, "",
            zrodlo="Zmienia KN o rzad wielkosci. Wartosc opisowa — nie wchodzi do formul.",
            podstawa="art. 5 ust. 7-8 — katalog nieprzytoczony w specyfikacji",
            format_liczby=TEKST, zolte=True)
    wejscie("sw_metoda_rz", "Metoda rozsadnego zysku",
            pzz.metoda_rozsadnego_zysku.value, "",
            zrodlo="Brak wzoru w specyfikacji. Wartosc opisowa.",
            podstawa="§ 6 ust. 5 rozp. 1897; § 12 ust. 10 rozp. 766",
            format_liczby=TEKST, zolte=True)
    wejscie("sw_prog_dwa", "Prog tolerancji przy dotacji i kredycie naraz",
            pzz.prog_tolerancji_przy_dwoch_instrumentach.value, "",
            zrodlo="Domyslnie 'nizszy' (10%). ZALOZENIE — odwraca werdykt testu 3. "
                   "Wartosc opisowa; prog liczbowy jest w zakladce Rekompensata.",
            podstawa="§ 7 ust. 9 rozp. 1897; § 13 ust. 8 rozp. 766 — zbieg nierozstrzygniety",
            format_liczby=TEKST, zolte=True)
    wejscie("sw_okres_rozliczeniowy", "Okres rozliczeniowy nadwyzki (lata)",
            pzz.okres_rozliczeniowy_nadwyzki_lat, "lata",
            zrodlo="Domyslnie 1 — kontrola w kazdym roku, wariant najostrozniejszy. "
                   "Dluzszy okres daje rzadsza siatke i lagodniejszy wynik. ZALOZENIE.",
            podstawa="§ 7 ust. 9 rozp. 1897; § 13 ust. 8 rozp. 766",
            format_liczby="0", zolte=True)
    wejscie("sw_bonus_prog", "Bonus +5 pp podnosi takze prog gruntowy",
            pzz.bonus_podnosi_prog_gruntowy, "TRUE/FALSE",
            zrodlo="Domyslnie FALSE — bonus podnosi wylacznie limit gorny. ZALOZENIE.",
            podstawa="art. 13 ust. 1 pkt 1 w zw. z art. 13 ust. 4 ustawy z 8.12.2006",
            format_liczby=TEKST, zolte=True)
    wejscie("sw_lokale_z_puli", "Lokale dla gminy pochodza z puli",
            pzz.lokale_dla_gminy_z_puli.value, "",
            zrodlo="Domyslnie proporcjonalnie z obu, kluczem PUM. Ustawa nie wskazuje "
                   "puli. ZALOZENIE. Wartosc opisowa.",
            podstawa="ustawa z 16.12.2020, Dz.U. 2021 poz. 223",
            format_liczby=TEKST, zolte=True)

    # --- katalog zalozen ---
    wiersz += 1
    wiersz = _sekcja(
        ws, wiersz,
        "KATALOG ZALOZEN — PELNA LISTA PYTAN DO BGK (patrz PYTANIA_DO_BGK.md)",
    )
    for kol, tytul in enumerate(
        ["Zalozenie", "Model przyjmuje", "Odczyt alternatywny", "Gdzie to zmienic"],
        start=1,
    ):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True, size=9)
        komorka.border = RAMKA_DOL
    wiersz += 1
    for z in _katalog_zalozen.wedlug_priorytetu():
        ws.cell(row=wiersz, column=1, value=f"{z.priorytet}. {z.tytul}").font = Font(size=9)
        ws.cell(row=wiersz, column=2, value=z.domyslnie).font = Font(size=9)
        ws.cell(row=wiersz, column=3, value=z.alternatywa).font = Font(size=9)
        ws.cell(
            row=wiersz, column=4,
            value=z.sciezka or "brak przelacznika — zmiana wymaga wejscia w kod",
        ).font = Font(size=9, color="FF666666")
        wiersz += 1

    # --- zastrzezenia ---
    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "ZASTRZEZENIA")
    for tekst in ZASTRZEZENIA:
        ws.cell(row=wiersz, column=1, value=tekst).font = Font(size=9, color="FF884400")
        ws.merge_cells(start_row=wiersz, start_column=1, end_row=wiersz, end_column=5)
        wiersz += 1

    ws.freeze_panes = "A5"


def _zrodlo_parametru(pz, nazwa: str) -> str:
    return "UZUPELNIC — parametr zewnetrzny bez wskazanego zrodla."


ZASTRZEZENIA = (
    "Narzedzie wspiera odsiewanie wariantow przed zlozeniem wniosku i rozmowe z gmina "
    "o proporcji pul. Nie zastepuje wyliczenia BGK, opinii prawnej ani doradztwa podatkowego.",
    "Rekompensate ustala wiazaco Bank. Wynik testu 3 z tego arkusza jest wskazaniem, nie rozstrzygnieciem.",
    "Stale prawne odczytano z tekstow jednolitych obowiazujacych w sierpniu 2026 r. Przed kazdym "
    "naborem sprawdz aktualnosc: parametry programow zmieniaja sie miedzy edycjami, a stopy "
    "zewnetrzne w ogole nie sa elementem prawa.",
    "Kwestie otwarte z rozdz. 10 specyfikacji pozostaja nierozstrzygniete i maja wplyw na wynik. "
    "Przelaczniki w sekcji powyzej pokazuja, ktore wartosci sa zalozeniami.",
    "Narzedzie nie ocenia zdolnosci kredytowej inwestora ani nie prowadzi analizy popytu — "
    "obloznosc jest parametrem wejsciowym.",
)


# ===========================================================================
# Zakladka Podstawy_prawne — wykaz stalych i tablice progowe do INDEX/MATCH
# ===========================================================================

def _podstawy_prawne(wb: Workbook, rej: Rejestr) -> None:
    ws = wb.create_sheet("Podstawy_prawne")
    _szerokosci(ws, [("A", 58), ("B", 16), ("C", 74)])

    wiersz = _naglowek(
        ws, 1, "Podstawy prawne stalych",
        "Kazda liczba uzyta w skoroszycie ma tu podstawe. Komorki z tej zakladki sa "
        "zrodlem dla formul w pozostalych — zmiana stalej przelicza caly model.",
    )
    for kol, tytul in enumerate(["Stala", "Wartosc", "Podstawa"], start=1):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True)
        komorka.border = RAMKA_DOL
    wiersz += 1

    # Stale wykorzystywane w formulach — kazda dostaje adres w rejestrze.
    stale = (
        ("grant_limit", "Grant spoleczny — limit podstawowy",
         prawo.GRANT_SPOLECZNY_LIMIT_PODSTAWOWY, "art. 13 ust. 1 pkt 1 ustawy z 8.12.2006"),
        ("grant_prog_gruntowy", "Grant spoleczny — prog gruntowy",
         prawo.GRANT_SPOLECZNY_PROG_GRUNTOWY, "art. 13 ust. 1 pkt 1 ustawy z 8.12.2006"),
        ("grant_komunalny", "Grant komunalny", prawo.GRANT_KOMUNALNY,
         "art. 13 ust. 1 pkt 3 lit. c w zw. z art. 5a ust. 1 ustawy z 8.12.2006"),
        ("bonus_pp", "Bonus rewitalizacyjny / Za zyciem", prawo.BONUS_REWITALIZACYJNY_PP,
         "art. 13 ust. 4 ustawy z 8.12.2006 — wylaczony przy finansowaniu zwrotnym"),
        ("hybryda_limit", "Limit laczny przy hybrydzie jako jednym przedsiewzieciu",
         prawo.GRANT_HYBRYDA_LIMIT_LACZNY, "art. 13 ust. 1a ustawy z 8.12.2006"),
        ("kredyt_max_udzial", "Kredyt — maksymalny udzial", prawo.KREDYT_MAKSYMALNY_UDZIAL,
         "art. 15b ust. 2 ustawy z 26.10.1995"),
        ("kredyt_max_okres", "Kredyt — maksymalny okres (lata)",
         Decimal(prawo.KREDYT_MAKSYMALNY_OKRES_LAT), "art. 15b ust. 3 ustawy z 26.10.1995"),
        ("limit_czynszu_remont", "Limit czynszu — remont i przebudowa",
         prawo.LIMIT_CZYNSZU_REMONT_I_PRZEBUDOWA,
         "art. 7c ustawy z 8.12.2006 w zw. z art. 5 ust. 1 pkt 2"),
        ("limit_czynszu_fz", "Limit czynszu — finansowanie zwrotne",
         prawo.LIMIT_CZYNSZU_FINANSOWANIE_ZWROTNE, "art. 28 ust. 2 pkt 2 ustawy z 26.10.1995"),
        ("limit_oplat", "Limit oplat poza czynszem", prawo.LIMIT_OPLAT_POZA_CZYNSZEM,
         "art. 28 ust. 4-5 ustawy z 26.10.1995 — osobny strumien, nie doliczany do czynszu"),
        ("part_prog_10", "Partycypacja — prog umowy bezterminowej",
         prawo.PARTYCYPACJA_PROG_UMOWA_BEZTERMINOWA, "art. 29a ust. 2b ustawy z 26.10.1995"),
        ("part_prog_15", "Partycypacja — prog wylaczenia art. 7b ust. 1",
         prawo.PARTYCYPACJA_PROG_WYLACZENIA_ART_7B, "art. 29a ust. 2a ustawy z 26.10.1995"),
        ("part_max", "Partycypacja — maksimum", prawo.PARTYCYPACJA_MAKSIMUM,
         "art. 29a ust. 2 ustawy z 26.10.1995"),
        ("okres_grant", "Okres powierzenia — grant (lata)",
         Decimal(prawo.OKRES_POWIERZENIA_GRANT_LAT), "art. 5 ust. 10 pkt 1 ustawy z 8.12.2006"),
        ("tol_grant", "Prog tolerancji nadwyzki — grant",
         prawo.PROG_TOLERANCJI_NADWYZKI_GRANT, "§ 7 ust. 9 rozp. Dz.U. 2025 poz. 1897"),
        ("tol_kredyt", "Prog tolerancji nadwyzki — kredyt",
         prawo.PROG_TOLERANCJI_NADWYZKI_KREDYT, "§ 13 ust. 8 rozp. t.j. Dz.U. 2021 poz. 766"),
        ("grunt_aport_limit", "Grunt z aportu w kosztach — sciezka kredytowa",
         prawo.GRUNT_APORT_LIMIT_W_KOSZTACH_KREDYT, "§ 12 ust. 7 rozp. t.j. Dz.U. 2021 poz. 766"),
        ("pum_min", "PUM lokalu — minimum (m2)", prawo.PUM_LOKALU_MIN_M2,
         "rozp. MIiR z 4.03.2019, Dz.U. 2019 poz. 457"),
        ("pum_max", "PUM lokalu — maksimum (m2)", prawo.PUM_LOKALU_MAX_M2,
         "rozp. MIiR z 4.03.2019 — powyzej wylacznie dla rodzin wielodzietnych"),
        ("dzwig_od", "Dzwigi osobowe — od kondygnacji",
         Decimal(prawo.DZWIG_OBOWIAZKOWY_OD_KONDYGNACJI), "rozp. MIiR z 4.03.2019"),
        ("droga_min", "Minimalna szerokosc drogi publicznej (m)",
         prawo.DROGA_PUBLICZNA_MIN_SZEROKOSC_M, "rozp. MIiR z 4.03.2019"),
    )
    for klucz, nazwa, wartosc, podstawa in stale:
        ws.cell(row=wiersz, column=1, value=nazwa)
        komorka = ws.cell(row=wiersz, column=2, value=float(wartosc))
        komorka.font = NIEBIESKI
        komorka.number_format = PROCENT_DOKLADNY if wartosc < 1 else "0"
        ws.cell(row=wiersz, column=3, value=podstawa).font = Font(size=9, color="FF666666")
        rej.zapisz(f"prawo.{klucz}", "Podstawy_prawne", _bezwzgledny("B", wiersz))
        wiersz += 1

    # --- tablica progowa art. 7c, rosnaco — do INDEX/MATCH z typem 1 ---
    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "TABELA LIMITOW CZYNSZU — art. 7c ustawy z 8.12.2006")
    ws.cell(row=wiersz, column=1, value="Udzial wsparcia od (wlacznie)").font = Font(bold=True)
    ws.cell(row=wiersz, column=2, value="Limit roczny").font = Font(bold=True)
    ws.cell(row=wiersz, column=3, value="Uwaga").font = Font(bold=True)
    wiersz += 1

    # `prawo.LIMIT_CZYNSZU_ART_7C` jest malejaca — do MATCH(...;1) potrzebna rosnaca.
    progi = sorted(prawo.LIMIT_CZYNSZU_ART_7C, key=lambda para: para[0])
    pierwszy = wiersz
    for prog, limit in progi:
        ws.cell(row=wiersz, column=1, value=float(prog)).number_format = PROCENT
        ws.cell(row=wiersz, column=1).font = NIEBIESKI
        ws.cell(row=wiersz, column=2, value=float(limit)).number_format = PROCENT_DOKLADNY
        ws.cell(row=wiersz, column=2).font = NIEBIESKI
        ws.cell(row=wiersz, column=3, value="progi czytane jako 'co najmniej'").font = Font(
            size=9, color="FF666666")
        wiersz += 1
    ostatni = wiersz - 1
    rej.zapisz("tabela.progi", "Podstawy_prawne", f"$A${pierwszy}:$A${ostatni}")
    rej.zapisz("tabela.limity", "Podstawy_prawne", f"$B${pierwszy}:$B${ostatni}")

    wiersz += 1
    ws.cell(row=wiersz, column=1,
            value="Limit z art. 7c i limit z art. 28 ust. 2 pkt 2 sa niezalezne — wiaze nizszy.").font = Font(
        italic=True, size=9, color="FF666666")
    wiersz += 2

    wiersz = _matryca_gruntu(ws, wiersz)
    ws.freeze_panes = "A5"


def _matryca_gruntu(ws: Worksheet, wiersz: int) -> int:
    """Matryca skutkow gruntu i zalozenia z rozdz. 9 uzupelnienia nr 2.

    Kazdy skutek niesie oznaczenie zrodla: Z — odczytane w przepisie, W — wniosek
    z odczytanych przepisow, ? — wymaga potwierdzenia w BGK. Bez tego czytelnik
    arkusza nie odroznilby przepisu od odczytu.
    """
    wiersz = _sekcja(ws, wiersz, "GRUNT — MATRYCA SKUTKOW (pochodzenie x forma)")
    ws.cell(
        row=wiersz, column=1,
        value="Z = odczytane w przepisie   |   W = wniosek z przepisow   |   "
              "? = do potwierdzenia w BGK",
    ).font = Font(italic=True, size=9, color="FF666666")
    wiersz += 1

    naglowki = (
        "Forma wniesienia", "Pochodzenie", "A: pasmo", "B: w kosztach",
        "C: przychod UOIG", "D: wydatek", "E: gmina wspolnikiem", "Podstawa",
    )
    for kol, tytul in enumerate(naglowki, start=1):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True)
        komorka.border = RAMKA_DOL
    wiersz += 1

    def znacznik(wartosc: bool, pewnosc: str, dopisek: str = "") -> str:
        return f"{'tak' if wartosc else 'nie'}{dopisek} [{pewnosc}]"

    for skutki in prawo.MATRYCA_GRUNTU:
        ws.cell(row=wiersz, column=1, value=skutki.forma)
        ws.cell(row=wiersz, column=2, value=skutki.pochodzenie)
        ws.cell(row=wiersz, column=3,
                value=znacznik(skutki.pasmo_45, skutki.pasmo_45_pewnosc))
        ws.cell(
            row=wiersz, column=4,
            value=znacznik(
                skutki.wartosc_w_kosztach,
                skutki.wartosc_w_kosztach_pewnosc,
                ", limit 20% w sciezce kredytowej" if skutki.limit_aportowy else "",
            ),
        )
        opis_c = znacznik(skutki.przychod_uoig, skutki.przychod_uoig_pewnosc)
        if skutki.przelacznik_przychodu:
            opis_c += f" — przelacznik '{skutki.przelacznik_przychodu}'"
        ws.cell(row=wiersz, column=5, value=opis_c)
        ws.cell(row=wiersz, column=6, value=skutki.wydatek)
        ws.cell(row=wiersz, column=7, value="tak" if skutki.gmina_wspolnikiem else "nie")
        ws.cell(row=wiersz, column=8, value=skutki.podstawa).font = Font(
            size=9, color="FF666666")
        wiersz += 1

    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "ZALOZENIA WYMAGAJACE POTWIERDZENIA W BGK — rozdz. 9")
    for kol, tytul in enumerate(
        ("Kwestia", "Przelacznik", "Odczyt domyslny", "Na czym polega", "Podstawa"), start=1
    ):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True)
        komorka.border = RAMKA_DOL
    wiersz += 1

    for numer, przelacznik, domyslna, opis, podstawa in (
        prawo.ZALOZENIA_GRUNTOWE_DO_POTWIERDZENIA
    ):
        ws.cell(row=wiersz, column=1, value=numer).font = Font(bold=True)
        ws.cell(row=wiersz, column=2, value=przelacznik)
        komorka = ws.cell(row=wiersz, column=3, value="TAK" if domyslna else "NIE")
        komorka.fill = ZOLTE
        ws.cell(row=wiersz, column=4, value=opis).font = Font(size=9)
        ws.cell(row=wiersz, column=5, value=podstawa).font = Font(size=9, color="FF666666")
        wiersz += 1

    wiersz += 1
    ws.cell(
        row=wiersz, column=1,
        value="Wartosc domyslna kazdego z powyzszych jest ZALOZENIEM, nie rozstrzygnieciem "
              "przepisu. Zmiana przelacznika w zakladce Zalozenia przelicza caly model.",
    ).font = Font(italic=True, size=9, color="FF666666")
    return wiersz + 1


# ===========================================================================
# Zakladka Alokacja — podzial kosztow wspolnych kluczem PUM
# ===========================================================================

def _alokacja(wb: Workbook, wynik: Wynik, rej: Rejestr) -> None:
    ws = wb.create_sheet("Alokacja")
    _szerokosci(ws, [("A", 42), ("B", 20), ("C", 20), ("D", 20), ("E", 60)])

    wiersz = _naglowek(
        ws, 1, "Alokacja kosztow wspolnych",
        "Koszty wspolne dziela sie miedzy pule proporcjonalnie do PUM. Wartosc gruntu "
        "dzieli sie tym samym kluczem — przypisanie jej obu pulom w calosci byloby "
        "policzeniem jej dwa razy.",
    )
    for kol, tytul in enumerate(
        ["Pozycja", "Lacznie", "Pula spoleczna", "Pula komunalna", "Uwaga"], start=1
    ):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True)
        komorka.border = RAMKA_DOL
    wiersz += 1

    def formula(kol: int, tresc: str, fmt: str = KWOTA, zielony: bool = True) -> None:
        komorka = ws.cell(row=wiersz, column=kol, value=tresc)
        komorka.number_format = fmt
        komorka.font = ZIELONY if zielony else CZARNY

    def etykieta(tekst: str, uwaga: str = "") -> None:
        ws.cell(row=wiersz, column=1, value=tekst)
        if uwaga:
            ws.cell(row=wiersz, column=5, value=uwaga).font = Font(size=9, color="FF666666")

    # --- klucz alokacji ---
    wiersz = _sekcja(ws, wiersz, "KLUCZ ALOKACJI")
    etykieta("Udzial puli komunalnej")
    formula(2, f"={rej['udzial_komunalny']}", PROCENT)
    rej.zapisz("alok.udzial_kom", "Alokacja", _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Udzial puli spolecznej")
    formula(2, f"=1-{rej['alok.udzial_kom']}", PROCENT, zielony=False)
    rej.zapisz("alok.udzial_spol", "Alokacja", _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("PUM")
    formula(2, f"={rej['pum_laczne']}", POWIERZCHNIA)
    formula(3, f"={rej['pum_laczne']}*{rej['alok.udzial_spol']}", POWIERZCHNIA, zielony=False)
    formula(4, f"={rej['pum_laczne']}*{rej['alok.udzial_kom']}", POWIERZCHNIA, zielony=False)
    rej.zapisz("alok.pum_laczne", "Alokacja", _bezwzgledny("B", wiersz))
    rej.zapisz("alok.pum_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.pum_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("PUM przychodowe",
             "Powierzchnia przynoszaca czynsz SIM. W trybie 'lokal za grunt' mniejsza od "
             "wybudowanej — lokale gminy trzeba wybudowac, ale nie przyniosa czynszu.")
    formula(2, f"={rej['alok.pum_laczne']}-{rej['grunt_pum_gminy']}", POWIERZCHNIA, zielony=False)
    formula(3, f"=$B${wiersz}*{rej['alok.udzial_spol']}", POWIERZCHNIA, zielony=False)
    formula(4, f"=$B${wiersz}*{rej['alok.udzial_kom']}", POWIERZCHNIA, zielony=False)
    rej.zapisz("alok.pum_przych_laczne", "Alokacja", _bezwzgledny("B", wiersz))
    rej.zapisz("alok.pum_przych_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.pum_przych_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Srednie PUM lokalu", "Poza przedzialem 25-80 m2 — patrz zakladka Werdykty.")
    formula(2, f"={rej['alok.pum_laczne']}/{rej['liczba_lokali']}", POWIERZCHNIA)
    rej.zapisz("alok.srednie_pum", "Alokacja", _bezwzgledny("B", wiersz))
    wiersz += 1

    # --- pozycje kosztowe ---
    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "POZYCJE KOSZTOWE (netto)")

    # Kanal B: wartosc gruntu wchodzaca do kosztow, juz po limicie z § 12 ust. 7.
    # Limit jest samozwrotny (grunt jest skladnikiem kosztow), wiec rozwiazany
    # w postaci zamknietej: u = K_bez_gruntu * limit / (1 - limit). Naklada sie go
    # PER PULA, bo przepis dotyczy przedsiewziecia finansowanego zwrotnie, a przy
    # domyslnym odczycie hybrydy tylko pula spoleczna korzysta z kredytu.
    etykieta("Grunt uznany w kosztach",
             "Kanal B. Dzierzawa daje 0; aport w puli kredytowej — limit § 12 ust. 7.")

    def _grunt_puli(udzial: str, kredytowa: bool) -> str:
        bez_gruntu = (
            f"({rej['koszt_budowy_m2']}*{rej['alok.pum_laczne']}+{rej['infrastruktura']}"
            f"+{rej['projekt_i_nadzor']}+{rej['koszty_ogolne']}+{rej['rezerwa']}"
            f"+{rej['dzwigi']})*{udzial}"
        )
        pelny = f"{rej['grunt_wartosc']}*{udzial}"
        if not kredytowa:
            return f"IF({rej['grunt_w_kosztach']}=FALSE,0,{pelny})"
        limit = rej["prawo.grunt_aport_limit"]
        return (
            f"IF({rej['grunt_w_kosztach']}=FALSE,0,"
            f"IF(AND({rej['grunt_aport']}=TRUE,{rej['kredyt_udzial']}>0),"
            f"MIN({pelny},{bez_gruntu}*{limit}/(1-{limit})),{pelny}))"
        )

    formula(3, "=" + _grunt_puli(rej["alok.udzial_spol"], True), KWOTA, zielony=False)
    formula(4, "=" + _grunt_puli(rej["alok.udzial_kom"], False), KWOTA, zielony=False)
    formula(2, f"=$C${wiersz}+$D${wiersz}", KWOTA, zielony=False)
    rej.zapisz("alok.grunt_uznany_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.grunt_uznany_kom", "Alokacja", _bezwzgledny("D", wiersz))
    rej.zapisz("alok.grunt_uznany", "Alokacja", _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Koszt budowy", "Przypisany wprost wg wlasnego PUM kazdej puli.")
    formula(2, f"={rej['koszt_budowy_m2']}*{rej['alok.pum_laczne']}")
    formula(3, f"={rej['koszt_budowy_m2']}*{rej['alok.pum_spol']}", KWOTA, zielony=False)
    formula(4, f"={rej['koszt_budowy_m2']}*{rej['alok.pum_kom']}", KWOTA, zielony=False)
    pierwszy_koszt = wiersz
    wiersz += 1

    pozycje_wspolne = (
        ("grunt", "Grunt (kanal B, po limicie aportowym)", "",
         "Klucz PUM, limit aportowy nakladany osobno na kazda pule."),
        ("infrastruktura", "Infrastruktura", rej["infrastruktura"], ""),
        ("projekt", "Projekt i nadzor", rej["projekt_i_nadzor"], ""),
        ("ogolne", "Koszty ogolne", rej["koszty_ogolne"], ""),
        ("rezerwa", "Rezerwa", rej["rezerwa"], ""),
        ("dzwigi", "Dzwigi osobowe", rej["dzwigi"], ""),
    )
    for klucz, nazwa, zrodlo, uwaga in pozycje_wspolne:
        etykieta(nazwa, uwaga)
        if klucz == "grunt":
            # Grunt ma juz wartosci per pula — limit aportowy dotyka tylko puli
            # kredytowej, wiec podzial nie jest prostym kluczem PUM.
            formula(3, f"={rej['alok.grunt_uznany_spol']}", KWOTA, zielony=False)
            formula(4, f"={rej['alok.grunt_uznany_kom']}", KWOTA, zielony=False)
            formula(2, f"=$C${wiersz}+$D${wiersz}", KWOTA, zielony=False)
            rej.zapisz("alok.grunt_spol", "Alokacja", _bezwzgledny("C", wiersz))
            rej.zapisz("alok.grunt_kom", "Alokacja", _bezwzgledny("D", wiersz))
        else:
            formula(2, f"={zrodlo}")
            formula(3, f"=$B${wiersz}*{rej['alok.udzial_spol']}", KWOTA, zielony=False)
            formula(4, f"=$B${wiersz}*{rej['alok.udzial_kom']}", KWOTA, zielony=False)
        wiersz += 1
    ostatni_koszt = wiersz - 1

    etykieta("Razem netto")
    for kol in (2, 3, 4):
        litera = get_column_letter(kol)
        komorka = ws.cell(
            row=wiersz, column=kol,
            value=f"=SUM({litera}{pierwszy_koszt}:{litera}{ostatni_koszt})",
        )
        komorka.number_format = KWOTA
        komorka.font = Font(bold=True)
    rej.zapisz("alok.netto_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.netto_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    # --- ujecie VAT ---
    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "KOSZTY PRZEDSIEWZIECIA (podstawa grantu i kwota do sfinansowania)")

    etykieta("Mnoznik VAT",
             "VAT wchodzi do podstawy tylko przy braku prawa do odliczenia — art. 13 ust. 3.")
    formula(2, f"=IF({rej['vat_odliczalny']}=TRUE,1,1+{rej['stawka_vat']})", "0.000")
    rej.zapisz("alok.mnoznik_vat", "Alokacja", _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Koszty przedsiewziecia")
    formula(2, f"=$C${wiersz}+$D${wiersz}", KWOTA, zielony=False)
    formula(3, f"={rej['alok.netto_spol']}*{rej['alok.mnoznik_vat']}", KWOTA, zielony=False)
    formula(4, f"={rej['alok.netto_kom']}*{rej['alok.mnoznik_vat']}", KWOTA, zielony=False)
    rej.zapisz("alok.koszty_laczne", "Alokacja", _bezwzgledny("B", wiersz))
    rej.zapisz("alok.koszty_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.koszty_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Grunt w podstawie (z VAT)", "Do limitu gruntowego grantu z art. 13 ust. 1 pkt 1.")
    formula(3, f"={rej['alok.grunt_spol']}*{rej['alok.mnoznik_vat']}", KWOTA, zielony=False)
    formula(4, f"={rej['alok.grunt_kom']}*{rej['alok.mnoznik_vat']}", KWOTA, zielony=False)
    rej.zapisz("alok.grunt_podstawa_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.grunt_podstawa_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Grunt do pasma dotacji (z VAT)",
             "Kanal A. Wartosc prawa z art. 13 ust. 1 pkt 1 — bez limitu aportowego. "
             "0, gdy forma nie daje ani wlasnosci, ani uzytkowania wieczystego.")
    formula(3, f"=IF({rej['grunt_pasmo_45']}=FALSE,0,{rej['grunt_do_pasma']}"
               f"*{rej['alok.udzial_spol']}*{rej['alok.mnoznik_vat']})", KWOTA, zielony=False)
    formula(4, f"=IF({rej['grunt_pasmo_45']}=FALSE,0,{rej['grunt_do_pasma']}"
               f"*{rej['alok.udzial_kom']}*{rej['alok.mnoznik_vat']})", KWOTA, zielony=False)
    rej.zapisz("alok.grunt_pasmo_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.grunt_pasmo_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Wklad rzeczowy w gruncie",
             "Kanal D. Wartosc gruntu, ktora siedzi w kosztach, ale nie wymaga wylozenia "
             "gotowki — aport, dzialka juz w spolce, prawo od gminy, rozliczenie lokalami.")
    formula(3, f"=IF({rej['grunt_gotowka']}=TRUE,0,{rej['alok.grunt_podstawa_spol']})",
            KWOTA, zielony=False)
    formula(4, f"=IF({rej['grunt_gotowka']}=TRUE,0,{rej['alok.grunt_podstawa_kom']})",
            KWOTA, zielony=False)
    rej.zapisz("alok.rzeczowy_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.rzeczowy_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Koszty bez gruntu",
             "Grunt ma w rekompensacie wlasna regule, wiec nie wchodzi tez do nakladu.")
    formula(3, f"={rej['alok.koszty_spol']}-{rej['alok.grunt_podstawa_spol']}", KWOTA, zielony=False)
    formula(4, f"={rej['alok.koszty_kom']}-{rej['alok.grunt_podstawa_kom']}", KWOTA, zielony=False)
    rej.zapisz("alok.koszty_bez_gruntu_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.koszty_bez_gruntu_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Koszt przedsiewziecia na m2 PUM",
             "Podstawa alternatywna limitu czynszu — art. 28 ust. 2b.")
    formula(3, f"=IF({rej['alok.pum_spol']}=0,0,{rej['alok.koszty_spol']}/{rej['alok.pum_spol']})",
            KWOTA_GROSZE, zielony=False)
    formula(4, f"=IF({rej['alok.pum_kom']}=0,0,{rej['alok.koszty_kom']}/{rej['alok.pum_kom']})",
            KWOTA_GROSZE, zielony=False)
    rej.zapisz("alok.koszt_m2_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("alok.koszt_m2_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    _granty_w_alokacji(ws, wiersz, rej)
    ws.freeze_panes = "A5"


def _granty_w_alokacji(ws: Worksheet, wiersz: int, rej: Rejestr) -> int:
    """Blok wysokosci wsparcia — dopisywany do zakladki Alokacja.

    Grant wynika wprost z podzielonych kosztow i z przypisanej wartosci gruntu,
    wiec siedzi w warstwie wspolnej. Obie pule odwoluja sie tu formula, dzieki
    czemu limit hybrydowy z art. 13 ust. 1a liczy sie bez odwolan cyklicznych.
    """
    def etykieta(tekst: str, uwaga: str = "") -> None:
        ws.cell(row=wiersz, column=1, value=tekst)
        if uwaga:
            ws.cell(row=wiersz, column=5, value=uwaga).font = Font(size=9, color="FF666666")

    def formula(kol: int, tresc: str, fmt: str = KWOTA) -> None:
        komorka = ws.cell(row=wiersz, column=kol, value=tresc)
        komorka.number_format = fmt
        komorka.font = CZARNY

    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "WYSOKOSC WSPARCIA")

    etykieta("Bonus czynny — pula spoleczna",
             "art. 13 ust. 4 wylacza bonus przy finansowaniu zwrotnym.")
    formula(3, f"=IF(AND({rej['bonus_spoleczny']}=TRUE,{rej['kredyt_udzial']}=0),"
               f"{rej['prawo.bonus_pp']},0)", PROCENT)
    formula(4, f"=IF({rej['bonus_komunalny']}=TRUE,{rej['prawo.bonus_pp']},0)", PROCENT)
    rej.zapisz("grant.bonus_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("grant.bonus_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Stawka nominalna")
    formula(3, f"={rej['prawo.grant_limit']}+{rej['grant.bonus_spol']}", PROCENT)
    formula(4, f"={rej['prawo.grant_komunalny']}+{rej['grant.bonus_kom']}", PROCENT)
    rej.zapisz("grant.stawka_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("grant.stawka_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Grunt wliczany do limitu",
             "Kanal A — art. 13 ust. 1 pkt 1. Forma bez prawa wlasnosci i bez "
             "uzytkowania wieczystego daje 0, niezaleznie od wartosci dzialki.")
    formula(3, f"={rej['alok.grunt_pasmo_spol']}")
    rej.zapisz("grant.grunt_wliczany", "Alokacja", _bezwzgledny("C", wiersz))
    wiersz += 1

    etykieta("Limit gorny (stawka x koszty)")
    formula(3, f"={rej['alok.koszty_spol']}*{rej['grant.stawka_spol']}")
    formula(4, f"={rej['alok.koszty_kom']}*{rej['grant.stawka_kom']}")
    rej.zapisz("grant.limit_gorny_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("grant.limit_gorny_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Limit gruntowy (prog x koszty + grunt)",
             "art. 13 ust. 1 pkt 1 — czesc ponad prog tylko do wysokosci gruntu inwestora.")
    formula(3, f"={rej['alok.koszty_spol']}*{rej['prawo.grant_prog_gruntowy']}"
               f"+{rej['grant.grunt_wliczany']}")
    ws.cell(row=wiersz, column=4, value="nie dotyczy").font = Font(size=9, color="FF666666")
    rej.zapisz("grant.limit_gruntowy_spol", "Alokacja", _bezwzgledny("C", wiersz))
    wiersz += 1

    etykieta("Grant przed limitem hybrydowym")
    formula(3, f"=MIN({rej['grant.limit_gorny_spol']},{rej['grant.limit_gruntowy_spol']})")
    formula(4, f"={rej['grant.limit_gorny_kom']}")
    rej.zapisz("grant.wstepny_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("grant.wstepny_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Wspolczynnik limitu hybrydowego",
             "art. 13 ust. 1a — czynny tylko przy przelaczniku 'hybryda jako jedno przedsiewziecie'.")
    formula(2,
            f"=IF(OR({rej['sw_hybryda']}=FALSE,"
            f"({rej['grant.wstepny_spol']}+{rej['grant.wstepny_kom']})=0),1,"
            f"MIN(1,{rej['alok.koszty_laczne']}*{rej['prawo.hybryda_limit']}/"
            f"({rej['grant.wstepny_spol']}+{rej['grant.wstepny_kom']})))", "0.0000")
    rej.zapisz("grant.wsp_hybryda", "Alokacja", _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("GRANT")
    for kol, klucz in ((3, "grant.wstepny_spol"), (4, "grant.wstepny_kom")):
        komorka = ws.cell(row=wiersz, column=kol,
                          value=f"={rej[klucz]}*{rej['grant.wsp_hybryda']}")
        komorka.number_format = KWOTA
        komorka.font = Font(bold=True)
    rej.zapisz("grant.spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("grant.kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1

    etykieta("Udzial wsparcia w kosztach", "Wejscie do tabeli limitow czynszu z art. 7c.")
    formula(3, f"=IF({rej['alok.koszty_spol']}=0,0,{rej['grant.spol']}/{rej['alok.koszty_spol']})",
            PROCENT_DOKLADNY)
    formula(4, f"=IF({rej['alok.koszty_kom']}=0,0,{rej['grant.kom']}/{rej['alok.koszty_kom']})",
            PROCENT_DOKLADNY)
    rej.zapisz("grant.udzial_spol", "Alokacja", _bezwzgledny("C", wiersz))
    rej.zapisz("grant.udzial_kom", "Alokacja", _bezwzgledny("D", wiersz))
    wiersz += 1
    return wiersz


# ===========================================================================
# Zakladki Pula_spoleczna i Pula_komunalna
# ===========================================================================

def _pula(wb: Workbook, wynik: Wynik, rej: Rejestr, spoleczna: bool) -> None:
    nazwa = "Pula_spoleczna" if spoleczna else "Pula_komunalna"
    p = "spol" if spoleczna else "kom"
    ws = wb.create_sheet(nazwa)
    _szerokosci(ws, [("A", 34)] + [(get_column_letter(k), 16) for k in range(2, 14)])

    proj = wynik.projekcja.spoleczna if spoleczna else wynik.projekcja.komunalna
    lat = proj.okres_powierzenia_lat
    kredyt = spoleczna

    podtytul = (
        "Projekcja rok po rok przez okres powierzenia. Kazda komorka jest formula — "
        "zmiana zalozenia przelicza cala zakladke."
    )
    if not spoleczna:
        podtytul += (
            " Pozycji kredytowych tu nie ma: art. 5a ust. 3 ustawy z 8.12.2006 wyklucza "
            "finansowanie zwrotne w tej puli."
        )
    wiersz = _naglowek(ws, 1, f"{nazwa.replace('_', ' ')}", podtytul)

    def etykieta(tekst: str, uwaga: str = "") -> None:
        ws.cell(row=wiersz, column=1, value=tekst)
        if uwaga:
            ws.cell(row=wiersz, column=6, value=uwaga).font = Font(size=9, color="FF666666")

    def wart(kol: int, tresc: str, fmt: str = KWOTA, pogrubione: bool = False,
             zielony: bool = False) -> None:
        komorka = ws.cell(row=wiersz, column=kol, value=tresc)
        komorka.number_format = fmt
        komorka.font = Font(bold=True) if pogrubione else (ZIELONY if zielony else CZARNY)

    # ---------------- finansowanie ----------------
    wiersz = _sekcja(ws, wiersz, "STRUKTURA FINANSOWANIA")

    etykieta("Koszty przedsiewziecia")
    wart(2, f"={rej[f'alok.koszty_{p}']}", KWOTA, zielony=True)
    rej.zapisz(f"{p}.koszty", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Grant")
    wart(2, f"={rej[f'grant.{p}']}", KWOTA, zielony=True)
    rej.zapisz(f"{p}.grant", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    if kredyt:
        automatyczny = wynik.wejscie.przelaczniki.tryb_kredytu is TrybKredytu.AUTOMATYCZNY
        etykieta(
            "Kredyt SBC",
            "Maksymalny kredyt, ktory uniesie zakladany czynsz. Maksimum 80% kosztow."
            if automatyczny
            else "Kwota z udzialu docelowego. Maksimum 80% kosztow.",
        )
        if not automatyczny:
            wart(2, f"={rej[f'{p}.koszty']}*{rej['kredyt_udzial']}")
        # W trybie automatycznym formula powstaje dopiero po projekcji — potrzebuje
        # kolumny pulapow, ktora liczy sie z przychodow i kosztow biezacych.
        komorka_kredytu = _bezwzgledny("B", wiersz)
        rej.zapisz(f"{p}.kredyt", nazwa, komorka_kredytu)
        wiersz += 1

        etykieta("Rata na zlotowke kredytu po karencji",
                 "Annuita jednostkowa — rata jest liniowa wzgledem kwoty kredytu.")
        wart(2,
             f"=IF(({rej['kredyt_n']}-{rej['kredyt_t']})<=0,0,"
             f"IF({rej['kredyt_rp']}=0,1/({rej['kredyt_n']}-{rej['kredyt_t']}),"
             f"{rej['kredyt_rp']}*(1+{rej['kredyt_rp']})^({rej['kredyt_n']}-{rej['kredyt_t']})"
             f"/((1+{rej['kredyt_rp']})^({rej['kredyt_n']}-{rej['kredyt_t']})-1)))", "0.00000000")
        rej.zapisz(f"{p}.annuita_jednostkowa", nazwa, _bezwzgledny("B", wiersz))
        wiersz += 1

        etykieta("Partycypacja", "art. 29a ust. 2 ustawy z 26.10.1995.")
        wart(2, f"={rej[f'{p}.koszty']}*{rej['partycypacja_stawka']}")
        rej.zapisz(f"{p}.partycypacja", nazwa, _bezwzgledny("B", wiersz))
        wiersz += 1
    else:
        etykieta("Kredyt SBC", "Wykluczony konstrukcyjnie — art. 5a ust. 3 ustawy z 8.12.2006.")
        wart(2, "=0")
        rej.zapisz(f"{p}.kredyt", nazwa, _bezwzgledny("B", wiersz))
        wiersz += 1

        etykieta("Partycypacja", "Nie wystepuje — najemca jest gmina.")
        wart(2, "=0")
        rej.zapisz(f"{p}.partycypacja", nazwa, _bezwzgledny("B", wiersz))
        wiersz += 1

    etykieta("Wklad wlasny")
    wart(2, f"={rej[f'{p}.koszty']}-{rej[f'{p}.grant']}-{rej[f'{p}.kredyt']}"
            f"-{rej[f'{p}.partycypacja']}", KWOTA, pogrubione=True)
    rej.zapisz(f"{p}.wklad", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    # ---------------- limity czynszu ----------------
    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "LIMITY CZYNSZU — wiaze nizszy")

    etykieta("Podstawa: wartosc odtworzeniowa 1 m2")
    wart(2, f"={rej['wartosc_odtworzeniowa']}", KWOTA_GROSZE, zielony=True)
    rej.zapisz(f"{p}.podstawa_wo", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Podstawa: koszt budowy z wartoscia nieruchomosci")
    wart(2, f"={rej[f'alok.koszt_m2_{p}']}", KWOTA_GROSZE, zielony=True)
    rej.zapisz(f"{p}.podstawa_kb", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Podstawa wiazaca", "art. 28 ust. 2b ustawy z 26.10.1995 — wyzsza z dwoch.")
    wart(2, f"=MAX({rej[f'{p}.podstawa_wo']},{rej[f'{p}.podstawa_kb']})", KWOTA_GROSZE, pogrubione=True)
    rej.zapisz(f"{p}.podstawa", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Udzial wsparcia w kosztach")
    wart(2, f"={rej[f'grant.udzial_{p}']}", PROCENT_DOKLADNY, zielony=True)
    rej.zapisz(f"{p}.udzial_wsparcia", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Stawka roczna z art. 7c",
             "INDEX/MATCH po tabeli progowej z zakladki Podstawy_prawne.")
    wart(2,
         f"=IF({rej['sw_remont']}=TRUE,{rej['prawo.limit_czynszu_remont']},"
         f"INDEX({rej['tabela.limity']},MATCH({rej[f'{p}.udzial_wsparcia']},"
         f"{rej['tabela.progi']},1)))", PROCENT_DOKLADNY)
    rej.zapisz(f"{p}.stawka_7c", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Limit z art. 7c")
    wart(2, f"={rej[f'{p}.podstawa']}*{rej[f'{p}.stawka_7c']}/12", STAWKA)
    rej.zapisz(f"{p}.limit_7c", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Limit z art. 28 ust. 2 pkt 2",
             "Dotyczy wylacznie lokali wybudowanych przy finansowaniu zwrotnym.")
    if kredyt:
        wart(2, f"=IF({rej[f'{p}.kredyt']}=0,\"nie dotyczy\","
                f"{rej[f'{p}.podstawa']}*{rej['prawo.limit_czynszu_fz']}/12)", STAWKA)
    else:
        wart(2, '="nie dotyczy"', TEKST)
    rej.zapisz(f"{p}.limit_28", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("LIMIT WIAZACY")
    if kredyt:
        wart(2, f"=IF({rej[f'{p}.kredyt']}=0,{rej[f'{p}.limit_7c']},"
                f"MIN({rej[f'{p}.limit_7c']},{rej[f'{p}.podstawa']}*"
                f"{rej['prawo.limit_czynszu_fz']}/12))", STAWKA, pogrubione=True)
    else:
        wart(2, f"={rej[f'{p}.limit_7c']}", STAWKA, pogrubione=True)
    rej.zapisz(f"{p}.limit_wiazacy", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Czynsz zakladany")
    zrodlo_czynszu = rej["czynsz_spoleczny"] if spoleczna else rej["czynsz_komunalny"]
    wart(2, f"={zrodlo_czynszu}", STAWKA, zielony=True)
    rej.zapisz(f"{p}.czynsz", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Zapas do limitu", "Wartosc ujemna oznacza czynsz ponad limit — konfiguracja bezprawna.")
    wart(2, f"={rej[f'{p}.limit_wiazacy']}-{rej[f'{p}.czynsz']}", STAWKA)
    rej.zapisz(f"{p}.zapas", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Pulap oplat poza czynszem",
             "art. 28 ust. 4-5 — osobny strumien, NIGDY nie doliczany do czynszu.")
    wart(2, f"={rej[f'{p}.podstawa_wo']}*{rej['prawo.limit_oplat']}/12", STAWKA)
    rej.zapisz(f"{p}.oplaty", nazwa, _bezwzgledny("B", wiersz))
    wiersz += 1

    if kredyt:
        wiersz = _harmonogram(ws, wiersz, rej, p, wynik)
    wiersz = _projekcja_arkusz(ws, wiersz, rej, p, lat, kredyt, spoleczna, proj)

    if kredyt and wynik.wejscie.przelaczniki.tryb_kredytu is TrybKredytu.AUTOMATYCZNY:
        # Kwota kredytu to najmniejszy pulap ze wszystkich lat, sciety limitem
        # ustawowym i zaokraglony w dol do pelnych zlotych — tak samo jak w silniku.
        komorka = ws[komorka_kredytu.replace("$", "")]
        # Trzy ograniczenia: ile uniesie czynsz, ile pozwala ustawa i ile
        # kredytu w ogole potrzeba po dotacji i partycypacji.
        # Wklad rzeczowy w gruncie juz pokrywa czesc kosztow, wiec kredyt nie ma
        # czego za niego finansowac — bez tego odjecia wynikowy wklad gotowkowy
        # wychodzilby ujemny. Odejmuje sie wklad rzeczowy TEJ puli — kredyt
        # finansuje przedsiewziecie spoleczne, a nie komunalne.
        potrzebny = (
            f"({rej[f'{p}.koszty']}-{rej[f'grant.{p}']}"
            f"-{rej[f'{p}.koszty']}*{rej['partycypacja_stawka']}"
            f"-{rej[f'alok.rzeczowy_{p}']})"
        )
        komorka.value = (
            f"=ROUNDDOWN(MAX(0,MIN(MIN({rej[f'{p}.pulapy']}),"
            f"{rej[f'{p}.koszty']}*{rej['prawo.kredyt_max_udzial']},{potrzebny})),0)"
        )
        komorka.number_format = KWOTA
        komorka.font = Font(bold=True)

        # Ktora z trzech wielkosci wiaze (pakiet nr 2, rozdz. 4). Bez tego
        # zachowanie modelu — czynsz rosnie, kredyt stoi — wyglada na blad.
        # Kolejnosc rozstrzygania przy remisie taka sama jak w silniku.
        wiersz += 1
        wiersz = _sekcja(ws, wiersz, "CO OGRANICZA KWOTE KREDYTU")
        limit_ustawowy = f"{rej[f'{p}.koszty']}*{rej['prawo.kredyt_max_udzial']}"
        for tytul, wzor, uwaga in (
            ("Udzwig czynszowy",
             f"=MIN({rej[f'{p}.pulapy']})",
             "Najwiekszy kredyt, ktory uniesie ten czynsz w najgorszym roku."),
            ("Limit ustawowy", f"={limit_ustawowy}",
             "art. 15b ust. 2 ustawy z 26.10.1995."),
            ("Faktyczna potrzeba", f"={potrzebny}",
             "Koszty po dotacji, partycypacji i wkladzie rzeczowym."),
        ):
            ws.cell(row=wiersz, column=1, value=tytul)
            ws.cell(row=wiersz, column=4, value=uwaga).font = Font(
                size=9, color="FF666666")
            k2 = ws.cell(row=wiersz, column=2, value=wzor)
            k2.number_format = KWOTA
            wiersz += 1

        ws.cell(row=wiersz, column=1, value="WIAZE").font = Font(bold=True)
        k2 = ws.cell(
            row=wiersz, column=2,
            value=(
                f'=IF({potrzebny}<=MIN(MIN({rej[f"{p}.pulapy"]}),{limit_ustawowy}),'
                f'"faktyczna potrzeba",'
                f'IF({limit_ustawowy}<=MIN({rej[f"{p}.pulapy"]}),'
                f'"limit ustawowy","udzwig czynszowy"))'
            ),
        )
        k2.number_format = TEKST
        k2.font = Font(bold=True)
        rej.zapisz(f"{p}.wiazace_kredyt", nazwa, _bezwzgledny("B", wiersz))
        wiersz += 1

    ws.freeze_panes = "B5"


def _harmonogram(ws: Worksheet, wiersz: int, rej: Rejestr, p: str, wynik: Wynik) -> int:
    """Harmonogram splat rownej raty z karencja splaty kapitalu."""
    lat = wynik.wejscie.pula_spoleczna.kredyt.okres_lat

    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "HARMONOGRAM SPLAT KREDYTU SBC")

    ws.cell(row=wiersz, column=1, value="Okresy splaty kapitalu (N - T)")
    komorka = ws.cell(row=wiersz, column=2, value=f"={rej['kredyt_n']}-{rej['kredyt_t']}")
    komorka.number_format = "0"
    rej.zapisz(f"{p}.okresy_splaty", "Pula_spoleczna", _bezwzgledny("B", wiersz))
    wiersz += 1

    ws.cell(row=wiersz, column=1, value="Rata annuitetowa po karencji")
    komorka = ws.cell(
        row=wiersz, column=2,
        value=(
            f"=IF({rej[f'{p}.okresy_splaty']}<=0,0,"
            f"IF({rej['kredyt_rp']}=0,{rej[f'{p}.kredyt']}/{rej[f'{p}.okresy_splaty']},"
            f"{rej[f'{p}.kredyt']}*{rej['kredyt_rp']}*(1+{rej['kredyt_rp']})^"
            f"{rej[f'{p}.okresy_splaty']}/((1+{rej['kredyt_rp']})^"
            f"{rej[f'{p}.okresy_splaty']}-1)))"
        ),
    )
    komorka.number_format = KWOTA
    komorka.font = Font(bold=True)
    rej.zapisz(f"{p}.rata", "Pula_spoleczna", _bezwzgledny("B", wiersz))
    wiersz += 1

    naglowki = ["Rok", "Saldo poczatkowe", "Odsetki", "Kapital", "Rata", "Saldo koncowe", "Karencja"]
    for kol, tytul in enumerate(naglowki, start=1):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True)
        komorka.border = RAMKA_DOL
    wiersz += 1

    pierwszy = wiersz
    for rok in range(1, lat + 1):
        w = wiersz
        komorka = ws.cell(row=w, column=1, value=str(rok))
        komorka.number_format = TEKST
        komorka.alignment = Alignment(horizontal="right")

        if rok == 1:
            saldo = f"={rej[f'{p}.kredyt']}"
        else:
            saldo = f"=F{w - 1}"
        ws.cell(row=w, column=2, value=saldo).number_format = KWOTA

        ws.cell(row=w, column=3, value=f"=B{w}*{rej['kredyt_rp']}").number_format = KWOTA
        ws.cell(
            row=w, column=4,
            value=(
                f"=IF({rok}<={rej['kredyt_t']},0,"
                f"IF({rok}>={rej['kredyt_n']},B{w},MAX(0,{rej[f'{p}.rata']}-C{w})))"
            ),
        ).number_format = KWOTA
        ws.cell(row=w, column=5, value=f"=C{w}+D{w}").number_format = KWOTA
        ws.cell(row=w, column=6, value=f"=B{w}-D{w}").number_format = KWOTA
        ws.cell(
            row=w, column=7,
            value=f'=IF({rok}<={rej["kredyt_t"]},"karencja","splata")',
        ).font = Font(size=9, color="FF666666")
        rej.zapisz(f"{p}.rata_rok_{rok}", "Pula_spoleczna", f"$E${w}")
        rej.zapisz(f"{p}.odsetki_rok_{rok}", "Pula_spoleczna", f"$C${w}")
        rej.zapisz_wiersz(f"{p}.harm_rok_{rok}", w)
        wiersz += 1

    ws.cell(row=wiersz, column=1, value="Razem").font = Font(bold=True)
    for kol in (3, 4, 5):
        litera = get_column_letter(kol)
        komorka = ws.cell(
            row=wiersz, column=kol, value=f"=SUM({litera}{pierwszy}:{litera}{wiersz - 1})"
        )
        komorka.number_format = KWOTA
        komorka.font = Font(bold=True)
    ws.cell(row=wiersz, column=6,
            value="Kapital razem musi rownac sie kwocie kredytu.").font = Font(
        size=9, color="FF666666")
    wiersz += 1
    return wiersz


def _projekcja_arkusz(
    ws: Worksheet, wiersz: int, rej: Rejestr, p: str, lat: int,
    kredyt: bool, spoleczna: bool, proj,
) -> int:
    """Przeplywy rok po rok przez okres powierzenia."""
    nazwa = "Pula_spoleczna" if spoleczna else "Pula_komunalna"

    wiersz += 1
    tytul = f"PROJEKCJA — OKRES POWIERZENIA {lat} LAT"
    tytul += (
        " (rowny okresowi finansowania, nie dluzej niz okres amortyzacji — § 11 rozp. 766)"
        if proj.sciezka == "kredyt"
        else " (art. 5 ust. 10 pkt 1 ustawy z 8.12.2006)"
    )
    wiersz = _sekcja(ws, wiersz, tytul)

    naglowki = [
        "Rok", "Indeks czynszu", "Indeks kosztow", "Czynsz zl/m2/mies.",
        "Przychod potencjalny", "Pustostany", "Przychod netto",
        "Eksploatacja", "Odpis remontowy", "Ubezpieczenie", "Zarzad", "Oplata za grunt",
        "Obsluga dlugu", "Rezerwa partycypacji", "Wymagane pokrycie",
        "Pokrycie wyplywow (test 2)", "Saldo",
        "Pulap kredytu", "Pokrycie obslugi dlugu (bankowe)",
    ]
    for kol, tytul_kol in enumerate(naglowki, start=1):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul_kol)
        komorka.font = Font(bold=True, size=9)
        komorka.border = RAMKA_DOL
        komorka.alignment = Alignment(wrap_text=True, vertical="bottom")
    wiersz += 1

    pustostany = (
        rej["pustostany"]
        if spoleczna
        else f'IF({rej["sw_pustostany_kom"]}=TRUE,{rej["pustostany"]},0)'
    )
    # Czynsz i koszty utrzymania licza sie od powierzchni przychodowej: lokale
    # oddane gminie w trybie "lokal za grunt" nie sa juz lokalami SIM. Ubezpieczenie
    # i zarzad zostaja na kluczu PUM, bo sa kosztem spolki, nie lokalu.
    pum = rej[f"alok.pum_przych_{p}"]
    udzial = rej[f"alok.udzial_{'spol' if spoleczna else 'kom'}"]
    # Oplata roczna za grunt obciaza koszty biezace tylko przy formach, ktore
    # rozliczaja sie oplata — dzierzawa i uzytkowanie wieczyste.
    oplata_gruntowa = (
        f'IF({rej["grunt_oplata_roczna"]}=0,0,{rej["grunt_oplata_roczna"]}*{udzial})'
    )

    pierwszy = wiersz
    for rok in range(1, lat + 1):
        w = wiersz
        komorka = ws.cell(row=w, column=1, value=str(rok))
        komorka.number_format = TEKST
        komorka.alignment = Alignment(horizontal="right")

        ws.cell(row=w, column=2, value=f"=(1+{rej['indeks_czynszu']})^({rok}-1)").number_format = "0.0000"
        ws.cell(row=w, column=3, value=f"=(1+{rej['indeks_kosztow']})^({rok}-1)").number_format = "0.0000"
        ws.cell(row=w, column=4, value=f"={rej[f'{p}.czynsz']}*B{w}").number_format = STAWKA
        ws.cell(row=w, column=5, value=f"=D{w}*12*{pum}").number_format = KWOTA
        ws.cell(row=w, column=6, value=f"=E{w}*{pustostany}").number_format = KWOTA
        ws.cell(row=w, column=7, value=f"=E{w}-F{w}").number_format = KWOTA
        ws.cell(row=w, column=8, value=f"={rej['koszt_ekspl_m2']}*{pum}*C{w}").number_format = KWOTA
        ws.cell(row=w, column=9, value=f"={rej['odpis_remontowy_m2']}*{pum}*C{w}").number_format = KWOTA
        ws.cell(row=w, column=10, value=f"={rej['ubezpieczenie']}*{udzial}*C{w}").number_format = KWOTA
        ws.cell(row=w, column=11, value=f"={rej['zarzad']}*{udzial}*C{w}").number_format = KWOTA
        ws.cell(row=w, column=12, value=f"={oplata_gruntowa}*C{w}").number_format = KWOTA

        if kredyt and rej.ma(f"{p}.rata_rok_{rok}"):
            ws.cell(row=w, column=13, value=f"={rej[f'{p}.rata_rok_{rok}']}").number_format = KWOTA
        else:
            ws.cell(row=w, column=13, value="=0").number_format = KWOTA

        if spoleczna:
            # art. 29a ust. 3 — zobowiazanie rosnie wskaznikiem GUS niezaleznie
            # od ponownego zasiedlenia.
            ws.cell(
                row=w, column=14,
                value=(
                    f"={rej[f'{p}.partycypacja']}*{rej['partycypacja_rotacja']}"
                    f"*(1+{rej['waloryzacja_part']})^{rok}"
                ),
            ).number_format = KWOTA
        else:
            ws.cell(row=w, column=14, value="=0").number_format = KWOTA

        ws.cell(row=w, column=15, value=f"=SUM(H{w}:M{w})").number_format = KWOTA
        ws.cell(
            row=w, column=16, value=f'=IF(O{w}=0,"",G{w}/O{w})'
        ).number_format = WSKAZNIK
        ws.cell(row=w, column=17, value=f"=G{w}-O{w}-N{w}").number_format = KWOTA
        if kredyt and rej.ma(f"{p}.annuita_jednostkowa"):
            # Rata jest liniowa wzgledem kwoty kredytu, wiec warunek pokrycia
            # w tym roku sprowadza sie do gornego pulapu kwoty.
            wspolczynnik = (
                f"IF({rok}<={rej['kredyt_t']},{rej['kredyt_rp']},"
                f"{rej[f'{p}.annuita_jednostkowa']})"
            )
            ws.cell(
                row=w, column=18,
                value=(f"=IF({rok}>{rej['kredyt_n']},\"\",IF({wspolczynnik}<=0,\"\","
                       f"MAX(0,(G{w}-SUM(H{w}:L{w}))/{rej['bufor_dscr']}"
                       f"/{wspolczynnik})))"),
            ).number_format = KWOTA
            # Kolumna S — pokrycie bankowe: nadwyzka operacyjna do samej raty.
            # To ta wielkosc ustawia bufor, a nie kolumna P (tam mianownikiem sa
            # wszystkie wyplywy). Dwie rozne liczby, dwie rozne role.
            ws.cell(
                row=w, column=19,
                value=f'=IF(M{w}<=0,"",(G{w}-SUM(H{w}:L{w}))/M{w})',
            ).number_format = WSKAZNIK
        rej.zapisz_wiersz(f"{p}.proj_rok_{rok}", w)
        wiersz += 1
    ostatni = wiersz - 1
    rej.zapisz_wiersz(f"{p}.proj_od", pierwszy)
    rej.zapisz_wiersz(f"{p}.proj_do", ostatni)

    ws.cell(row=wiersz, column=1, value="Razem").font = Font(bold=True)
    for kol in (5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 17):
        litera = get_column_letter(kol)
        komorka = ws.cell(
            row=wiersz, column=kol, value=f"=SUM({litera}{pierwszy}:{litera}{ostatni})"
        )
        komorka.number_format = KWOTA
        komorka.font = Font(bold=True)
    wiersz += 1

    ws.cell(row=wiersz, column=1, value="Minimalne pokrycie wyplywow (test 2)").font = Font(bold=True)
    komorka = ws.cell(row=wiersz, column=16, value=f"=MIN(P{pierwszy}:P{ostatni})")
    komorka.number_format = WSKAZNIK
    komorka.font = Font(bold=True)
    rej.zapisz(f"{p}.min_dscr", nazwa, _bezwzgledny("P", wiersz))
    wiersz += 1

    # Rejestrowane zawsze, takze w puli komunalnej — tam kredyt jest niedopuszczalny
    # (art. 5a ust. 3), wiec kolumna S jest pusta i komorka pokazuje "brak dlugu".
    # Pusty rejestr wywrocilby tabele werdyktow, ktora czyta oba klucze.
    ws.cell(
        row=wiersz, column=1, value="Minimalne pokrycie obslugi dlugu (bankowe)"
    ).font = Font(bold=True)
    komorka = ws.cell(
        row=wiersz, column=19,
        value=f'=IF(COUNT(S{pierwszy}:S{ostatni})=0,"brak dlugu",MIN(S{pierwszy}:S{ostatni}))',
    )
    komorka.number_format = WSKAZNIK
    komorka.font = Font(bold=True)
    rej.zapisz(f"{p}.min_pokrycie_dlugu", nazwa, _bezwzgledny("S", wiersz))
    wiersz += 1

    if kredyt and rej.ma(f"{p}.annuita_jednostkowa"):
        rej.zapisz(f"{p}.pulapy", nazwa, f"$R${pierwszy}:$R${ostatni}")

    ws.cell(row=wiersz, column=1, value="Lat z pokryciem wyplywow ponizej 1,0").font = Font(bold=True)
    komorka = ws.cell(
        row=wiersz, column=16,
        value=f'=COUNTIF(P{pierwszy}:P{ostatni},"<1")',
    )
    komorka.number_format = "0"
    komorka.font = Font(bold=True)
    rej.zapisz(f"{p}.lat_naruszenia", nazwa, _bezwzgledny("P", wiersz))
    wiersz += 1
    return wiersz


# ===========================================================================
# Zakladka Rekompensata
# ===========================================================================

def _rekompensata(wb: Workbook, wynik: Wynik, rej: Rejestr) -> None:
    ws = wb.create_sheet("Rekompensata")
    _szerokosci(ws, [("A", 40)] + [(get_column_letter(k), 17) for k in range(2, 12)])

    wiersz = _naglowek(
        ws, 1, "Rekompensata — KN, RZ, EDB i test nadwyzki",
        "Warunek graniczny: RUOIG <= KN + RZ (art. 5 ust. 5 i 11 ustawy z 8.12.2006; "
        "§ 6 ust. 2 rozp. Dz.U. 2025 poz. 1897). Sciezki liczone osobno, chyba ze "
        "przelacznik hybrydy mowi inaczej. WYLICZENIE BANKU JEST WIAZACE.",
    )

    for spoleczna in (True, False):
        p = "spol" if spoleczna else "kom"
        proj = wynik.projekcja.spoleczna if spoleczna else wynik.projekcja.komunalna
        rek = wynik.rekompensata.spoleczna if spoleczna else wynik.rekompensata.komunalna
        if rek is None:
            continue
        wiersz = _sekcja(
            ws, wiersz,
            f"PULA {'SPOLECZNA' if spoleczna else 'KOMUNALNA'} — sciezka {proj.sciezka.upper()}, "
            f"okres powierzenia {rek.okres_powierzenia_lat} lat",
        )
        wiersz = _rekompensata_puli(ws, wiersz, rej, wynik, p, spoleczna, rek, proj)
        wiersz += 1

    ws.freeze_panes = "A5"


def _rekompensata_puli(ws, wiersz, rej, wynik, p, spoleczna, rek, proj):
    lat = rek.okres_powierzenia_lat
    sciezka_kredytowa = proj.sciezka == "kredyt"

    def etykieta(tekst: str, uwaga: str = "") -> None:
        ws.cell(row=wiersz, column=1, value=tekst)
        if uwaga:
            ws.cell(row=wiersz, column=4, value=uwaga).font = Font(size=9, color="FF666666")

    # --- ujecie gruntu ---
    etykieta("Ujecie gruntu", rek.grunt_ujecie)
    wiersz += 1

    # Kanaly B i C wykluczaja sie: grunt, ktorego inwestor nie kupil, nie jest
    # kosztem uslugi, a grunt, ktory kupil, nie jest jego przychodem. Wartosc
    # w kosztach jest juz po limicie z § 12 ust. 7, bo limit naklada Alokacja.
    etykieta("Grunt jako koszt (rok 1)",
             "Kanal B. 0, gdy wartosc gruntu jest przychodem uslugi publicznej.")
    wzor = (
        f"=IF({rej['grunt_przychod_uoig']}=TRUE,0,{rej[f'alok.grunt_podstawa_{p}']})"
    )
    komorka = ws.cell(row=wiersz, column=2, value=wzor)
    komorka.number_format = KWOTA
    rej.zapisz(f"rek.{p}.grunt_koszt", "Rekompensata", _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Grunt jako przychod (rok 1)",
             "Kanal C — art. 5 ust. 9 pkt 4. Grunt wniesiony przez gmine obniza KN. "
             "Przychod bierze pelna wartosc z operatu, bez limitu aportowego.")
    wzor = (
        f"=IF({rej['grunt_przychod_uoig']}=TRUE,{rej['grunt_wartosc']}"
        f"*{rej['alok.udzial_' + p]},0)"
    )
    komorka = ws.cell(row=wiersz, column=2, value=wzor)
    komorka.number_format = KWOTA
    rej.zapisz(f"rek.{p}.grunt_przychod", "Rekompensata", _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("Roczny naklad inwestycyjny w KUOIG",
             f"Ujecie '{wynik.wejscie.przelaczniki.koszty_inwestycyjne_w_kn.value}'. "
             "Bez gruntu — ten ma wlasna regule. ZALOZENIE, patrz LUKI.md.")
    ujecie = wynik.wejscie.przelaczniki.koszty_inwestycyjne_w_kn
    baza = rej[f"alok.koszty_bez_gruntu_{p}"]
    if ujecie is UjecieKosztowInwestycyjnych.POMINIETE:
        wzor = "=0"
    elif ujecie is UjecieKosztowInwestycyjnych.NAKLAD_POCZATKOWY:
        wzor = f"={baza}"
    elif ujecie is UjecieKosztowInwestycyjnych.AMORTYZACJA_W_OKRESIE_POWIERZENIA:
        wzor = f"={baza}/{lat}"
    else:
        wzor = f"={baza}/{rej['okres_amortyzacji']}"
    komorka = ws.cell(row=wiersz, column=2, value=wzor)
    komorka.number_format = KWOTA
    rej.zapisz(f"rek.{p}.naklad", "Rekompensata", _bezwzgledny("B", wiersz))
    wiersz += 1

    tylko_rok_pierwszy = ujecie is UjecieKosztowInwestycyjnych.NAKLAD_POCZATKOWY

    # --- tabela lat ---
    wiersz += 1
    naglowki = ["Rok", "Koszty biezace", "Naklad", "Odsetki", "Grunt-koszt",
                "KUOIG", "Czynsz", "Grunt-przychod", "PUOIG", "Czynnik dyskonta", "Netto zdysk."]
    for kol, tytul in enumerate(naglowki, start=1):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True, size=9)
        komorka.border = RAMKA_DOL
        komorka.alignment = Alignment(wrap_text=True)
    wiersz += 1

    arkusz_puli = "Pula_spoleczna" if spoleczna else "Pula_komunalna"
    pierwszy = wiersz
    for rok in range(1, lat + 1):
        w = wiersz
        wiersz_puli = rej.wiersz(f"{p}.proj_rok_{rok}")
        komorka = ws.cell(row=w, column=1, value=str(rok))
        komorka.number_format = TEKST
        komorka.alignment = Alignment(horizontal="right")

        # art. 5 ust. 7-8 — koszty biezace. Koszty stale zarzadu sa w sciezce
        # kredytowej wskazane wprost przez § 12 ust. 3 rozp. 766. Oplata roczna
        # za grunt (kolumna L) obciaza usluge przez caly okres powierzenia.
        ws.cell(row=w, column=2,
                value=f"=SUM('{arkusz_puli}'!H{wiersz_puli}:L{wiersz_puli})").number_format = KWOTA
        ws.cell(row=w, column=3,
                value=(f"=IF({rok}=1,{rej[f'rek.{p}.naklad']},0)" if tylko_rok_pierwszy
                       else f"={rej[f'rek.{p}.naklad']}")).number_format = KWOTA
        odsetki = (
            f"={rej[f'{p}.odsetki_rok_{rok}']}"
            if sciezka_kredytowa and rej.ma(f"{p}.odsetki_rok_{rok}")
            else "=0"
        )
        ws.cell(row=w, column=4, value=odsetki).number_format = KWOTA
        ws.cell(row=w, column=5,
                value=f"=IF({rok}=1,{rej[f'rek.{p}.grunt_koszt']},0)").number_format = KWOTA
        ws.cell(row=w, column=6, value=f"=SUM(B{w}:E{w})").number_format = KWOTA
        ws.cell(row=w, column=7, value=f"='{arkusz_puli}'!G{wiersz_puli}").number_format = KWOTA
        ws.cell(row=w, column=8,
                value=f"=IF({rok}=1,{rej[f'rek.{p}.grunt_przychod']},0)").number_format = KWOTA
        ws.cell(row=w, column=9, value=f"=G{w}+H{w}").number_format = KWOTA
        ws.cell(row=w, column=10,
                value=f"=(1+{rej['rb']})^({rok}-1)").number_format = "0.0000"
        ws.cell(row=w, column=11, value=f"=(F{w}-I{w})/J{w}").number_format = KWOTA
        rej.zapisz(f"rek.{p}.netto_rok_{rok}", "Rekompensata", _bezwzgledny("K", w))
        wiersz += 1
    ostatni = wiersz - 1

    etykieta("KOSZTY NETTO (KN)", "Suma zdyskontowanych roznic KUOIG - PUOIG.")
    komorka = ws.cell(row=wiersz, column=11, value=f"=SUM(K{pierwszy}:K{ostatni})")
    komorka.number_format = KWOTA
    komorka.font = Font(bold=True)
    rej.zapisz(f"rek.{p}.kn", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    # --- rozsadny zysk ---
    wiersz += 1
    etykieta("Rozsadny zysk (RZ)",
             f"Metoda '{wynik.wejscie.przelaczniki.metoda_rozsadnego_zysku.value}'. "
             "Brak wzoru w specyfikacji — ZALOZENIE, patrz LUKI.md.")
    if wynik.wejscie.przelaczniki.metoda_rozsadnego_zysku is MetodaRozsadnegoZysku.KWOTA_WPROST:
        klucz_udzialu = "alok.udzial_kom" if p == "kom" else "alok.udzial_spol"
        wzor = f"={rej['rz_kwota']}*{rej[klucz_udzialu]}"
    else:
        # Zdyskontowany strumien godziwego zwrotu z kapitalu wlasnego, stopa IRS BGK.
        wzor = (
            f"=IF(MAX(0,{rej[f'{p}.wklad']})=0,0,MAX(0,{rej[f'{p}.wklad']})*{rej['irs']}"
            f"*IF({rej['rb']}=0,{lat},(1-(1+{rej['rb']})^(-{lat}))/{rej['rb']}*(1+{rej['rb']})))"
        )
    komorka = ws.cell(row=wiersz, column=11, value=wzor)
    komorka.number_format = KWOTA
    rej.zapisz(f"rek.{p}.rz", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    # --- EDB ---
    etykieta("EDB grantu", "§ 4 pkt 1 rozp. Dz.U. 2018 poz. 461 — rowny kwocie dotacji.")
    komorka = ws.cell(row=wiersz, column=11, value=f"={rej[f'{p}.grant']}")
    komorka.number_format = KWOTA
    rej.zapisz(f"rek.{p}.edb_grant", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    if sciezka_kredytowa:
        wiersz = _edb_kredytu(ws, wiersz, rej, p, wynik)
    else:
        etykieta("EDB kredytu", "Brak finansowania zwrotnego w tej sciezce.")
        komorka = ws.cell(row=wiersz, column=11, value="=0")
        komorka.number_format = KWOTA
        rej.zapisz(f"rek.{p}.edb_kredyt", "Rekompensata", _bezwzgledny("K", wiersz))
        wiersz += 1

    etykieta("Wsparcie dodatkowe (RFRM + dokumentacja BGK)",
             "§ 7 ust. 7 rozp. Dz.U. 2025 poz. 1897.")
    udzial = rej[f"alok.udzial_{'spol' if spoleczna else 'kom'}"]
    komorka = ws.cell(row=wiersz, column=11,
                      value=f"=({rej['rfrm']}+{rej['dokumentacja_bgk']})*{udzial}")
    komorka.number_format = KWOTA
    rej.zapisz(f"rek.{p}.dodatkowe", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    # --- test nadwyzki ---
    wiersz += 1
    etykieta("REKOMPENSATA (RUOIG)")
    komorka = ws.cell(
        row=wiersz, column=11,
        value=f"={rej[f'rek.{p}.edb_grant']}+{rej[f'rek.{p}.edb_kredyt']}+{rej[f'rek.{p}.dodatkowe']}",
    )
    komorka.number_format = KWOTA
    komorka.font = Font(bold=True)
    rej.zapisz(f"rek.{p}.ruoig", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    etykieta("Dopuszczalna (KN + RZ)")
    komorka = ws.cell(row=wiersz, column=11, value=f"={rej[f'rek.{p}.kn']}+{rej[f'rek.{p}.rz']}")
    komorka.number_format = KWOTA
    komorka.font = Font(bold=True)
    rej.zapisz(f"rek.{p}.dopuszczalna", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    etykieta("Nadwyzka")
    komorka = ws.cell(
        row=wiersz, column=11,
        value=f"=MAX(0,{rej[f'rek.{p}.ruoig']}-{rej[f'rek.{p}.dopuszczalna']})",
    )
    komorka.number_format = KWOTA
    rej.zapisz(f"rek.{p}.nadwyzka", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    etykieta("Srednia roczna rekompensata")
    komorka = ws.cell(row=wiersz, column=11, value=f"={rej[f'rek.{p}.ruoig']}/{lat}")
    komorka.number_format = KWOTA
    rej.zapisz(f"rek.{p}.srednia", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    pula_wyniku = _pula_rekompensaty(wynik, p)
    prog_wartosc = (
        pula_wyniku.prog_tolerancji if pula_wyniku is not None
        else prawo.prog_tolerancji_nadwyzki("kredyt" if sciezka_kredytowa else "grant")
    )
    alternatywny = pula_wyniku.prog_tolerancji_alternatywny if pula_wyniku else None
    if alternatywny is None:
        uwaga_progu = "§ 13 ust. 8 rozp. 766" if sciezka_kredytowa else "§ 7 ust. 9 rozp. 1897"
    else:
        uwaga_progu = (
            "ZBIEG PROGOW: pula ma dotacje i kredyt naraz, a zaden przepis nie mowi, "
            f"ktory rezim wiaze. Przyjeto {prog_wartosc:.0%}; odczyt alternatywny to "
            f"{alternatywny:.0%}. Komorka jest ZALOZENIEM — podmien i przelicz, zeby "
            "zobaczyc drugi wariant. Patrz LUKI.md."
        )
    etykieta("Prog tolerancji nadwyzki — udzial", uwaga_progu)
    komorka = ws.cell(row=wiersz, column=11, value=float(prog_wartosc))
    komorka.number_format = PROCENT_DOKLADNY
    if alternatywny is not None:
        komorka.fill = ZOLTE
    rej.zapisz(f"rek.{p}.prog", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    wiersz = _profil_nadwyzki(ws, wiersz, rej, p, wynik, lat, sciezka_kredytowa)

    etykieta("Prog tolerancji nadwyzki — kwotowo",
             "Udzial progu odniesiony do sredniej rocznej rekompensaty.")
    komorka = ws.cell(
        row=wiersz, column=11,
        value=f"={rej[f'rek.{p}.srednia']}*{rej[f'rek.{p}.prog']}",
    )
    komorka.number_format = KWOTA
    rej.zapisz(f"rek.{p}.tolerancja", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    etykieta("Kwota do zwrotu do Funduszu Doplat")
    komorka = ws.cell(
        row=wiersz, column=11,
        value=f"=IF({rej[f'rek.{p}.najgorsza']}<={rej[f'rek.{p}.prog']},0,"
              f"{rej[f'rek.{p}.nadwyzka']})",
    )
    komorka.number_format = KWOTA
    komorka.font = CZERWONY
    rej.zapisz(f"rek.{p}.zwrot", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    etykieta("WERDYKT", "Na roku najgorszym profilu, nie na sredniej okresu.")
    komorka = ws.cell(
        row=wiersz, column=11,
        value=f'=IF({rej[f"rek.{p}.najgorsza"]}<={rej[f"rek.{p}.prog"]},'
              f'"przechodzi","nie przechodzi")',
    )
    komorka.font = Font(bold=True)
    rej.zapisz(f"rek.{p}.werdykt", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1
    return wiersz


def _pula_rekompensaty(wynik: Wynik, p: str):
    """Pula rekompensaty odpowiadajaca skrotowi 'spol'/'kom' w rejestrze."""
    return wynik.rekompensata.spoleczna if p == "spol" else wynik.rekompensata.komunalna


def _profil_nadwyzki(
    ws: Worksheet, wiersz: int, rej: Rejestr, p: str, wynik: Wynik,
    lat: int, sciezka_kredytowa: bool,
) -> int:
    """Profil nadwyzki rok po roku — pakiet naprawczy nr 2, rozdz. 3.

    Rachunek usredniony zakladal rowny rozklad. Tu rekompensata i kwota
    dopuszczalna narastaja rok po roku, a werdykt bierze rok najgorszy. Wszystkie
    skladniki maja rozklad wynikajacy z modelu, wiec arkusz liczy je formulami —
    nie przepisuje wyniku silnika.
    """
    metoda = wynik.wejscie.przelaczniki.metoda_rozsadnego_zysku
    wiersz += 1
    ws.cell(row=wiersz, column=1, value="PROFIL NADWYZKI — narastajaco").font = Font(bold=True)
    ws.cell(
        row=wiersz, column=4,
        value="Werdykt na roku najgorszym. Dotacja splywa na poczatku, "
              "przychody czynszowe przez caly okres.",
    ).font = Font(size=9, color="FF666666")
    wiersz += 1

    naglowki = ["Rok", "RUOIG rok", "RUOIG narast.", "RZ rok", "Dopuszczalna rok",
                "Dopuszczalna narast.", "Nadwyzka narast.", "Nadwyzka wzgledna"]
    for kol, tytul in enumerate(naglowki, start=1):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True, size=9)
        komorka.border = RAMKA_DOL
        komorka.alignment = Alignment(wrap_text=True)
    wiersz += 1

    pierwszy = wiersz
    for rok in range(1, lat + 1):
        w = wiersz
        komorka = ws.cell(row=w, column=1, value=str(rok))
        komorka.number_format = TEKST
        komorka.alignment = Alignment(horizontal="right")

        # EDB grantu i wsparcie dodatkowe splywaja na etapie inwestycji — rok 1.
        skladniki = []
        if rok == 1:
            skladniki.append(f"{rej[f'rek.{p}.edb_grant']}+{rej[f'rek.{p}.dodatkowe']}")
        if sciezka_kredytowa and rej.ma(f"rek.{p}.edb_rok_{rok}"):
            skladniki.append(rej[f"rek.{p}.edb_rok_{rok}"])
        ws.cell(row=w, column=2, value="=" + ("+".join(skladniki) or "0")).number_format = KWOTA
        ws.cell(
            row=w, column=3,
            value=f"=B{w}" if rok == 1 else f"=C{w - 1}+B{w}",
        ).number_format = KWOTA

        if metoda is MetodaRozsadnegoZysku.KWOTA_WPROST:
            # Kwota podana wprost nie ma wlasnego profilu — rozklad rowny jest
            # DODATKOWYM zalozeniem, tym samym co w silniku.
            rz_rok = f"={rej[f'rek.{p}.rz']}/{lat}"
        else:
            rz_rok = (
                f"=MAX(0,{rej[f'{p}.wklad']})*{rej['irs']}/(1+{rej['rb']})^({rok}-1)"
            )
        ws.cell(row=w, column=4, value=rz_rok).number_format = KWOTA
        ws.cell(
            row=w, column=5,
            value=f"={rej[f'rek.{p}.netto_rok_{rok}']}+D{w}",
        ).number_format = KWOTA
        ws.cell(
            row=w, column=6,
            value=f"=E{w}" if rok == 1 else f"=F{w - 1}+E{w}",
        ).number_format = KWOTA
        ws.cell(row=w, column=7, value=f"=MAX(0,C{w}-F{w})").number_format = KWOTA
        ws.cell(
            row=w, column=8, value=f'=IF(C{w}<=0,0,G{w}/C{w})'
        ).number_format = PROCENT_DOKLADNY
        wiersz += 1
    ostatni = wiersz - 1

    ws.cell(
        row=wiersz, column=1, value="Najgorszy rok — nadwyzka wzgledna"
    ).font = Font(bold=True)
    komorka = ws.cell(row=wiersz, column=11, value=f"=MAX(H{pierwszy}:H{ostatni})")
    komorka.number_format = PROCENT_DOKLADNY
    komorka.font = Font(bold=True)
    rej.zapisz(f"rek.{p}.najgorsza", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    ws.cell(row=wiersz, column=1, value="Najgorszy rok — numer").font = Font(bold=True)
    komorka = ws.cell(
        row=wiersz, column=11,
        value=f"=INDEX(A{pierwszy}:A{ostatni},MATCH(MAX(H{pierwszy}:H{ostatni}),"
              f"H{pierwszy}:H{ostatni},0))",
    )
    komorka.number_format = TEKST
    rej.zapisz(f"rek.{p}.rok_najgorszy", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1
    return wiersz


def _edb_kredytu(ws: Worksheet, wiersz: int, rej: Rejestr, p: str, wynik: Wynik) -> int:
    """EDB kredytu wg § 4 pkt 5 lit. e rozp. Dz.U. 2018 poz. 461 — rozpisane po latach."""
    k = wynik.wejscie.pula_spoleczna.kredyt
    N = k.okres_lat

    ws.cell(row=wiersz, column=1, value="EDB kredytu — § 4 pkt 5 lit. e").font = Font(bold=True)
    ws.cell(row=wiersz, column=4,
            value="Asercja: 0 <= EDB < S. Wynik ujemny oznacza rp > r.").font = Font(
        size=9, color="FF666666")
    wiersz += 1

    ws.cell(row=wiersz, column=1, value="Annuita przy stopie referencyjnej r")
    komorka = ws.cell(
        row=wiersz, column=2,
        value=(
            f"=IF({rej[f'{p}.okresy_splaty']}<=0,0,IF({rej['r']}=0,"
            f"{rej[f'{p}.kredyt']}/{rej[f'{p}.okresy_splaty']},"
            f"{rej[f'{p}.kredyt']}*{rej['r']}*(1+{rej['r']})^{rej[f'{p}.okresy_splaty']}"
            f"/((1+{rej['r']})^{rej[f'{p}.okresy_splaty']}-1)))"
        ),
    )
    komorka.number_format = KWOTA
    rej.zapisz("edb.annuita_r", "Rekompensata", _bezwzgledny("B", wiersz))
    wiersz += 1

    ws.cell(row=wiersz, column=1, value="Annuita przy stopie preferencyjnej rp")
    komorka = ws.cell(
        row=wiersz, column=2,
        value=(
            f"=IF({rej[f'{p}.okresy_splaty']}<=0,0,IF({rej['kredyt_rp']}=0,"
            f"{rej[f'{p}.kredyt']}/{rej[f'{p}.okresy_splaty']},"
            f"{rej[f'{p}.kredyt']}*{rej['kredyt_rp']}*(1+{rej['kredyt_rp']})^"
            f"{rej[f'{p}.okresy_splaty']}/((1+{rej['kredyt_rp']})^"
            f"{rej[f'{p}.okresy_splaty']}-1)))"
        ),
    )
    komorka.number_format = KWOTA
    rej.zapisz("edb.annuita_rp", "Rekompensata", _bezwzgledny("B", wiersz))
    wiersz += 1

    for kol, tytul in enumerate(["Okres i", "Czlon karencji", "Czlon splaty", "Dyskonto",
                                 "Wklad do EDB"], start=1):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True, size=9)
        komorka.border = RAMKA_DOL
    wiersz += 1

    pierwszy = wiersz
    for i in range(1, N + 1):
        w = wiersz
        komorka = ws.cell(row=w, column=1, value=str(i))
        komorka.number_format = TEKST
        komorka.alignment = Alignment(horizontal="right")
        # Czlon pierwszy dziala tylko w karencji: S*r - S*rp.
        ws.cell(
            row=w, column=2,
            value=(f"=IF({i}<={rej['kredyt_t']},"
                   f"{rej[f'{p}.kredyt']}*{rej['r']}-{rej[f'{p}.kredyt']}*{rej['kredyt_rp']},0)"),
        ).number_format = KWOTA
        # Czlon drugi dziala po karencji: roznica annuit.
        ws.cell(
            row=w, column=3,
            value=f"=IF({i}>{rej['kredyt_t']},{rej['edb.annuita_r']}-{rej['edb.annuita_rp']},0)",
        ).number_format = KWOTA
        ws.cell(row=w, column=4, value=f"=(1+{rej['rd']})^{i}").number_format = "0.0000"
        ws.cell(row=w, column=5, value=f"=(B{w}+C{w})/D{w}").number_format = KWOTA
        # Wklad roku i do EDB — potrzebny profilowi nadwyzki (pakiet nr 2, rozdz. 3).
        rej.zapisz(f"rek.{p}.edb_rok_{i}", "Rekompensata", _bezwzgledny("E", w))
        wiersz += 1
    ostatni = wiersz - 1

    ws.cell(row=wiersz, column=1, value="EDB kredytu").font = Font(bold=True)
    komorka = ws.cell(row=wiersz, column=11, value=f"=SUM(E{pierwszy}:E{ostatni})")
    komorka.number_format = KWOTA
    komorka.font = Font(bold=True)
    rej.zapisz(f"rek.{p}.edb_kredyt", "Rekompensata", _bezwzgledny("K", wiersz))
    wiersz += 1

    ws.cell(row=wiersz, column=1, value="Kontrola asercji 0 <= EDB < S")
    komorka = ws.cell(
        row=wiersz, column=11,
        value=(f'=IF({rej[f"rek.{p}.edb_kredyt"]}<0,"BLAD: rp > r",'
               f'IF({rej[f"rek.{p}.edb_kredyt"]}>={rej[f"{p}.kredyt"]},'
               f'"BLAD: EDB >= S","w normie"))'),
    )
    komorka.font = Font(bold=True)
    wiersz += 1
    return wiersz


# ===========================================================================
# Zakladka Werdykty
# ===========================================================================

def _werdykty(wb: Workbook, wynik: Wynik, rej: Rejestr) -> None:
    ws = wb.create_sheet("Werdykty")
    _szerokosci(ws, [("A", 44), ("B", 22), ("C", 22), ("D", 76)])

    wiersz = _naglowek(
        ws, 1, "Trzy testy montazu",
        "Projekt domyka sie WYLACZNIE gdy przechodza wszystkie trzy. Kazdy werdykt "
        "negatywny podaje wiazace ograniczenie i luke w liczbach.",
    )

    def etykieta(tekst: str, uwaga: str = "") -> None:
        ws.cell(row=wiersz, column=1, value=tekst)
        if uwaga:
            ws.cell(row=wiersz, column=4, value=uwaga).font = Font(size=9, color="FF666666")

    def wart(tresc: str, fmt: str = KWOTA, kol: int = 2, pogrubione: bool = False) -> None:
        komorka = ws.cell(row=wiersz, column=kol, value=tresc)
        komorka.number_format = fmt
        if pogrubione:
            komorka.font = Font(bold=True)

    # --- test 1 ---
    wiersz = _sekcja(ws, wiersz, "TEST 1 — KAPITAL (wklad wlasny jest WYNIKIEM, nie wejsciem)")
    etykieta("Koszty przedsiewziecia")
    wart(f"={rej['alok.koszty_laczne']}")
    wiersz += 1
    etykieta("Grant")
    wart(f"={rej['grant.spol']}+{rej['grant.kom']}")
    wiersz += 1
    etykieta("Kredyt SBC")
    wart(f"={rej['spol.kredyt']}+{rej['kom.kredyt']}")
    wiersz += 1
    etykieta("Partycypacja")
    wart(f"={rej['spol.partycypacja']}+{rej['kom.partycypacja']}")
    wiersz += 1
    etykieta("Wklad rzeczowy w gruncie",
             "Kanal D. Grunt wniesiony aportem albo prawo ustanowione przez gmine siedzi "
             "w kosztach, ale nikt za nie nie placi gotowka. Nabycie daje 0.")
    wart(f"={rej['alok.rzeczowy_spol']}+{rej['alok.rzeczowy_kom']}", KWOTA)
    rej.zapisz("werd.wklad_rzeczowy", "Werdykty", _bezwzgledny("B", wiersz))
    wiersz += 1
    etykieta("WYMAGANY WKLAD WLASNY",
             "Glowna liczba wyjsciowa narzedzia. Jeden bilans inwestora — obie pule "
             "skladaja sie na to samo zapotrzebowanie. Test pyta o pieniadze, wiec "
             "wklad rzeczowy w gruncie jest odjety.")
    wart(
        f"={rej['spol.wklad']}+{rej['kom.wklad']}-{rej['werd.wklad_rzeczowy']}",
        KWOTA,
        pogrubione=True,
    )
    rej.zapisz("werd.wklad_wymagany", "Werdykty", _bezwzgledny("B", wiersz))
    wiersz += 1
    etykieta("Udzial wkladu w kosztach")
    wart(f"={rej['werd.wklad_wymagany']}/{rej['alok.koszty_laczne']}", PROCENT)
    wiersz += 1
    etykieta("Wymagany wklad na m2 PUM")
    wart(f"={rej['werd.wklad_wymagany']}/{rej['alok.pum_laczne']}", KWOTA_GROSZE)
    wiersz += 1
    etykieta("Zadeklarowany kapital inwestora",
             "Opcjonalny punkt odniesienia. Pusty — narzedzie podaje sama wymagana kwote.")
    wart(f'=IF({rej["wklad_dostepny"]}="","nie podano",{rej["wklad_dostepny"]})')
    wiersz += 1
    etykieta("BRAKUJACY KAPITAL")
    wart(f'=IF({rej["wklad_dostepny"]}="",0,'
         f'MAX(0,{rej["werd.wklad_wymagany"]}-{rej["wklad_dostepny"]}))',
         KWOTA, pogrubione=True)
    ws.cell(row=wiersz, column=2).font = CZERWONY
    rej.zapisz("werd.luka_kapitalowa", "Werdykty", _bezwzgledny("B", wiersz))
    wiersz += 1
    etykieta("WERDYKT 1")
    wart(f'=IF({rej["werd.luka_kapitalowa"]}=0,"przechodzi","nie przechodzi")',
         TEKST, pogrubione=True)
    rej.zapisz("werd.test1", "Werdykty", _bezwzgledny("B", wiersz))
    wiersz += 1

    # --- test 2 ---
    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "TEST 2 — ZDOLNOSC CZYNSZOWA")
    ws.cell(row=wiersz, column=2, value="Pula spoleczna").font = Font(bold=True)
    ws.cell(row=wiersz, column=3, value="Pula komunalna").font = Font(bold=True)
    wiersz += 1

    for nazwa, klucz_spol, klucz_kom, fmt in (
        ("Limit czynszu wiazacy", "spol.limit_wiazacy", "kom.limit_wiazacy", STAWKA),
        ("Czynsz zakladany", "spol.czynsz", "kom.czynsz", STAWKA),
        ("Zapas do limitu", "spol.zapas", "kom.zapas", STAWKA),
        ("Minimalne pokrycie wyplywow (test 2)", "spol.min_dscr", "kom.min_dscr", WSKAZNIK),
        ("Minimalne pokrycie obslugi dlugu (bankowe)",
         "spol.min_pokrycie_dlugu", "kom.min_pokrycie_dlugu", WSKAZNIK),
        ("Lat z pokryciem wyplywow ponizej 1,0",
         "spol.lat_naruszenia", "kom.lat_naruszenia", "0"),
    ):
        etykieta(nazwa)
        wart(f"={rej[klucz_spol]}", fmt, kol=2)
        wart(f"={rej[klucz_kom]}", fmt, kol=3)
        wiersz += 1

    etykieta("Czynsz rynkowy (pula spoleczna)",
             "Sufit faktyczny. Pusty = nie podano; silnik nie podstawia zadnej wartosci. "
             "W puli komunalnej nie wystepuje — najemca jest gmina.")
    wart(f'=IF({rej["czynsz_rynkowy"]}="","nie podano",{rej["czynsz_rynkowy"]})',
         STAWKA, kol=2)
    wiersz += 1

    etykieta("Sufit wiazacy (pula spoleczna)",
             "Nizszy z dwoch: limit ustawowy albo rynek. Limit prawny przesuwa sie "
             "zmiana udzialu dotacji; sufitu rynkowego nie przesunie nic.")
    wart(
        f'=IF({rej["czynsz_rynkowy"]}="","limit ustawowy",'
        f'IF({rej["czynsz_rynkowy"]}<{rej["spol.limit_wiazacy"]},"rynek","limit ustawowy"))',
        TEKST, kol=2,
    )
    wiersz += 1

    etykieta("Pulap oplat poza czynszem",
             "art. 28 ust. 4-5 — osobny strumien, nie doliczany do czynszu.")
    wart(f"={rej['spol.oplaty']}", STAWKA, kol=2)
    wart(f"={rej['kom.oplaty']}", STAWKA, kol=3)
    wiersz += 1

    etykieta("LUKA CZYNSZOWA — pula spoleczna",
             "O ile trzeba podniesc stawke bazowa, zeby DSCR nie schodzil ponizej 1,0.")
    wart(f"=IF({rej['spol.min_dscr']}>=1,0,{rej['spol.czynsz']}*(1/{rej['spol.min_dscr']}-1))",
         STAWKA, pogrubione=True)
    ws.cell(row=wiersz, column=2).font = CZERWONY
    rej.zapisz("werd.luka_czynszowa", "Werdykty", _bezwzgledny("B", wiersz))
    wiersz += 1

    etykieta("WERDYKT 2")
    wart(f'=IF(AND({rej["spol.lat_naruszenia"]}=0,{rej["kom.lat_naruszenia"]}=0),'
         f'"przechodzi","nie przechodzi")', TEKST, pogrubione=True)
    rej.zapisz("werd.test2", "Werdykty", _bezwzgledny("B", wiersz))
    wiersz += 1

    # --- test 3 ---
    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "TEST 3 — REKOMPENSATA")
    ws.cell(row=wiersz, column=4,
            value="WYLICZENIE BGK JEST WIAZACE. To jest wskazanie, nie rozstrzygniecie.").font = CZERWONY
    ws.cell(row=wiersz, column=2, value="Pula spoleczna").font = Font(bold=True)
    ws.cell(row=wiersz, column=3, value="Pula komunalna").font = Font(bold=True)
    wiersz += 1

    for nazwa, klucz, fmt in (
        ("Koszty netto (KN)", "kn", KWOTA),
        ("Rozsadny zysk (RZ)", "rz", KWOTA),
        ("Rekompensata (RUOIG)", "ruoig", KWOTA),
        ("Dopuszczalna (KN + RZ)", "dopuszczalna", KWOTA),
        ("Nadwyzka", "nadwyzka", KWOTA),
        ("Prog tolerancji", "tolerancja", KWOTA),
        ("Kwota do zwrotu", "zwrot", KWOTA),
        ("Werdykt czastkowy", "werdykt", TEKST),
    ):
        etykieta(nazwa)
        for kol, p in ((2, "spol"), (3, "kom")):
            if rej.ma(f"rek.{p}.{klucz}"):
                wart(f"={rej[f'rek.{p}.{klucz}']}", fmt, kol=kol)
            else:
                ws.cell(row=wiersz, column=kol, value="nie dotyczy").font = Font(
                    size=9, color="FF666666")
        wiersz += 1

    etykieta("NADWYZKA LACZNA")
    czesci = [f"{rej[f'rek.{p}.nadwyzka']}" for p in ("spol", "kom") if rej.ma(f"rek.{p}.nadwyzka")]
    wart("=" + "+".join(czesci) if czesci else "=0", KWOTA, pogrubione=True)
    ws.cell(row=wiersz, column=2).font = CZERWONY
    wiersz += 1

    etykieta("WERDYKT 3")
    warunki = [f'{rej[f"rek.{p}.werdykt"]}="przechodzi"' for p in ("spol", "kom")
               if rej.ma(f"rek.{p}.werdykt")]
    wart(f'=IF(AND({",".join(warunki)}),"przechodzi","nie przechodzi")' if warunki
         else '="przechodzi"', TEKST, pogrubione=True)
    rej.zapisz("werd.test3", "Werdykty", _bezwzgledny("B", wiersz))
    wiersz += 1

    # --- werdykt zbiorczy ---
    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "WERDYKT ZBIORCZY")
    etykieta("CZY MONTAZ SIE DOMYKA")
    wart(
        f'=IF(AND({rej["werd.test1"]}="przechodzi",{rej["werd.test2"]}="przechodzi",'
        f'{rej["werd.test3"]}="przechodzi"),"TAK","NIE")', TEKST, pogrubione=True
    )
    wiersz += 1

    etykieta("WIAZACE OGRANICZENIE")
    komorka = ws.cell(
        row=wiersz, column=2,
        value=(
            f'=IF({rej["werd.test1"]}<>"przechodzi","Test 1 — kapital: brakuje wkladu wlasnego",'
            f'IF({rej["werd.test2"]}<>"przechodzi","Test 2 — zdolnosc czynszowa: DSCR ponizej 1,0",'
            f'IF({rej["werd.test3"]}<>"przechodzi","Test 3 — rekompensata: nadwyzka ponad prog",'
            f'"brak — wszystkie trzy testy przechodza")))'
        ),
    )
    komorka.font = Font(bold=True)
    ws.merge_cells(start_row=wiersz, start_column=2, end_row=wiersz, end_column=4)
    wiersz += 1

    # --- ostrzezenia silnika ---
    wiersz += 2
    wiersz = _sekcja(ws, wiersz, "OSTRZEZENIA I ZALOZENIA SILNIKA")
    for o in wynik.ostrzezenia:
        ws.cell(row=wiersz, column=1, value=o.kod).font = Font(bold=True, size=9)
        komorka = ws.cell(row=wiersz, column=2, value=o.tresc)
        komorka.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=wiersz, start_column=2, end_row=wiersz, end_column=3)
        ws.cell(row=wiersz, column=4, value=o.podstawa).font = Font(size=9, color="FF666666")
        ws.row_dimensions[wiersz].height = 30
        wiersz += 1

    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "ZASTRZEZENIA")
    for tekst in ZASTRZEZENIA:
        ws.cell(row=wiersz, column=1, value=tekst).font = Font(size=9, color="FF884400")
        ws.merge_cells(start_row=wiersz, start_column=1, end_row=wiersz, end_column=4)
        wiersz += 1

    ws.freeze_panes = "A5"


# ===========================================================================
# Zakladka Wrazliwosc
# ===========================================================================

def _wrazliwosc(wb: Workbook, wynik: Wynik, analiza: Optional[Analiza]) -> None:
    ws = wb.create_sheet("Wrazliwosc")
    _szerokosci(ws, [("A", 30), ("B", 12), ("C", 12), ("D", 12), ("E", 16), ("F", 18), ("G", 70)])

    wiersz = _naglowek(
        ws, 1, "Wrazliwosc i punkt graniczny",
        "UWAGA: ta zakladka jest MIGAWKA z silnika, nie zywymi formulami. Kazdy punkt "
        "sweepu to osobne przeliczenie calego modelu, ktorego nie da sie zlozyc z formul "
        "jednej zakladki. Zywy model siedzi w zakladkach Alokacja, Pula_* i Rekompensata — "
        "po zmianie zalozenia wygeneruj arkusz ponownie, zeby odswiezyc te zakladke.",
    )

    if analiza is None:
        ws.cell(row=wiersz, column=1,
                value="Analiza wrazliwosci nie zostala policzona dla tego eksportu.").font = Font(
            italic=True, color="FF666666")
        return

    ws.cell(row=wiersz, column=1, value=analiza.podsumowanie).font = Font(bold=True, size=11)
    ws.merge_cells(start_row=wiersz, start_column=1, end_row=wiersz, end_column=7)
    wiersz += 2

    # --- sweep ---
    wiersz = _sekcja(ws, wiersz, "SWEEP UDZIALU PULI KOMUNALNEJ")
    naglowki = ["Udzial komunalny", "Test 1", "Test 2", "Test 3", "Domyka sie",
                "Luka", "Wiazace ograniczenie"]
    for kol, tytul in enumerate(naglowki, start=1):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True, size=9)
        komorka.border = RAMKA_DOL
    wiersz += 1

    for punkt in analiza.sweep.punkty:
        ws.cell(row=wiersz, column=1, value=float(punkt.udzial)).number_format = PROCENT
        if punkt.policzalny:
            for kol, zdany in enumerate(punkt.werdykty, start=2):
                komorka = ws.cell(row=wiersz, column=kol, value="tak" if zdany else "nie")
                komorka.font = ZIELONY_BOLD if zdany else CZERWONY
            komorka = ws.cell(row=wiersz, column=5,
                              value="DOMYKA" if punkt.domyka_sie else "nie domyka")
            komorka.font = ZIELONY_BOLD if punkt.domyka_sie else CZERWONY
            if punkt.luki:
                opis, kwota, jednostka = punkt.luki[0]
                ws.cell(row=wiersz, column=6, value=float(kwota)).number_format = (
                    KWOTA if jednostka == "zl" else STAWKA
                )
        else:
            for kol in range(2, 6):
                ws.cell(row=wiersz, column=kol, value="—").font = Font(color="FF999999")
        tresc = punkt.powod_niepoliczalnosci or punkt.wiazace_ograniczenie
        ws.cell(row=wiersz, column=7, value=tresc[:220]).font = Font(size=9, color="FF666666")
        wiersz += 1

    wiersz += 1
    for nazwa, wartosc in (
        ("Maksymalny udzial puli komunalnej, przy ktorym wszystko przechodzi",
         analiza.sweep.maksymalny_udzial_komunalny),
        ("Punkt graniczny — pierwszy udzial, przy ktorym przestaje sie domykac",
         analiza.sweep.punkt_graniczny),
    ):
        ws.cell(row=wiersz, column=1, value=nazwa).font = Font(bold=True)
        komorka = ws.cell(row=wiersz, column=5,
                          value=float(wartosc) if wartosc is not None else "brak")
        komorka.number_format = PROCENT if wartosc is not None else TEKST
        komorka.font = Font(bold=True)
        wiersz += 1

    if analiza.sweep.test_blokujacy is not None:
        ws.cell(row=wiersz, column=1, value="Test blokujacy w calym zakresie").font = Font(bold=True)
        ws.cell(row=wiersz, column=5, value=f"Test {analiza.sweep.test_blokujacy}").font = CZERWONY
        wiersz += 1

    # --- ranking ---
    wiersz += 1
    wiersz = _sekcja(ws, wiersz, "RANKING PARAMETROW — ktory najtaniej przesuwa punkt graniczny")
    naglowki = ["Parametr", "Bazowo", "-20%", "+20%", "Granica -20%", "Granica +20%",
                "Sila wplywu / dzwignia"]
    for kol, tytul in enumerate(naglowki, start=1):
        komorka = ws.cell(row=wiersz, column=kol, value=tytul)
        komorka.font = Font(bold=True, size=9)
        komorka.border = RAMKA_DOL
    wiersz += 1

    for r in analiza.ranking:
        ws.cell(row=wiersz, column=1, value=r.nazwa)
        ws.cell(row=wiersz, column=2, value=float(r.wartosc_bazowa)).number_format = "#,##0.0000"
        ws.cell(row=wiersz, column=3, value=float(r.wartosc_dol)).number_format = "#,##0.0000"
        ws.cell(row=wiersz, column=4, value=float(r.wartosc_gora)).number_format = "#,##0.0000"
        for kol, granica in ((5, r.maks_udzial_dol), (6, r.maks_udzial_gora)):
            komorka = ws.cell(row=wiersz, column=kol,
                              value=float(granica) if granica is not None else "brak")
            komorka.number_format = PROCENT if granica is not None else TEKST
        ws.cell(row=wiersz, column=7,
                value=f"{float(r.sila_wplywu):.0%} — {r.opis_dzwigni}")
        wiersz += 1

    wiersz += 1
    ws.cell(row=wiersz, column=1,
            value="Wrazliwosc na stope referencyjna KE liczona jest zawsze, niezaleznie od "
                  "konfiguracji: wynik testu rekompensaty jest na nia bardzo wrazliwy, "
                  "a horyzont siega 30 lat.").font = Font(italic=True, size=9, color="FF666666")
    ws.merge_cells(start_row=wiersz, start_column=1, end_row=wiersz, end_column=7)
    ws.freeze_panes = "A5"


# ===========================================================================
# Eksport
# ===========================================================================

def zbuduj_skoroszyt(wynik: Wynik, analiza: Optional[Analiza] = None) -> Workbook:
    """Buduje skoroszyt z formulami. Nic nie zapisuje."""
    wb = Workbook()
    wb.remove(wb.active)
    rej = Rejestr()

    # Kolejnosc budowania wynika z zaleznosci formul; kolejnosc zakladek
    # ustawiana jest na koncu wg rozdz. 8.2 specyfikacji.
    _zalozenia(wb, wynik, rej)
    _podstawy_prawne(wb, rej)
    _alokacja(wb, wynik, rej)
    _pula(wb, wynik, rej, spoleczna=True)
    _pula(wb, wynik, rej, spoleczna=False)
    _rekompensata(wb, wynik, rej)
    _werdykty(wb, wynik, rej)
    _wrazliwosc(wb, wynik, analiza)

    kolejnosc = [
        "Zalozenia", "Alokacja", "Pula_spoleczna", "Pula_komunalna",
        "Rekompensata", "Werdykty", "Wrazliwosc", "Podstawy_prawne",
    ]
    wb._sheets = [wb[nazwa] for nazwa in kolejnosc if nazwa in wb.sheetnames]
    return wb


def eksportuj(
    wynik: Wynik, sciezka: Path | str, analiza: Optional[Analiza] = None
) -> Path:
    """Buduje skoroszyt i zapisuje go pod wskazana sciezka."""
    sciezka = Path(sciezka)
    sciezka.parent.mkdir(parents=True, exist_ok=True)
    zbuduj_skoroszyt(wynik, analiza).save(sciezka)
    return sciezka


def eksportuj_do_strumienia(wynik: Wynik, strumien, analiza: Optional[Analiza] = None):
    """Zapisuje skoroszyt do strumienia bajtow.

    Potrzebne tam, gdzie nie ma zapisywalnego systemu plikow — na przyklad
    w funkcji serverless, gdzie arkusz wraca prosto do przegladarki.
    """
    zbuduj_skoroszyt(wynik, analiza).save(strumien)
    return strumien
