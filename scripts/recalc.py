#!/usr/bin/env python3
"""Przeliczenie arkusza i kontrola bledow formul.

Arkusz zawiera formuly, wiec przed wydaniem trzeba go OBOWIAZKOWO przeliczyc
i zazadac zera bledow. Zielony przebieg dowodzi, ze formuly sie licza — nie ze
sa poprawne. Dodatkowo sprawdz recznie 2-3 formuly w kazdej zakladce projekcji.

Uzycie:
    python3 scripts/recalc.py plik.xlsx [--porownaj-z-silnikiem wejscie.yaml]

Openpyxl zapisuje formuly bez zapamietanych wynikow, wiec LibreOffice musi je
policzyc przy wczytaniu — stad przeliczenie jest realne, a nie odczytem cache'u.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Dict, List, Tuple

import openpyxl

BLEDY_EXCEL = (
    "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "#ERROR",
    "Err:", "#ADRES!", "#ARG!", "#DZIEL/0!", "#NAZWA?", "#LICZBA!", "#PUSTY!",
)


def przelicz(zrodlo: Path) -> Path:
    """Przelicza skoroszyt LibreOffice i zwraca sciezke do wersji z wynikami."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise SystemExit(
            "Nie znaleziono LibreOffice (soffice). Bez niego nie da sie przeliczyc "
            "arkusza, a wydanie nieprzeliczonego arkusza jest niedopuszczalne."
        )
    katalog = Path(tempfile.mkdtemp(prefix="recalc_"))
    wynik = subprocess.run(
        [
            soffice, "--headless", "--norestore", "--nolockcheck",
            "--convert-to", "xlsx:Calc MS Excel 2007 XML",
            "--outdir", str(katalog), str(zrodlo),
        ],
        capture_output=True,
        text=True,
        timeout=600,
    )
    przeliczony = katalog / (zrodlo.stem + ".xlsx")
    if not przeliczony.exists():
        raise SystemExit(
            "LibreOffice nie wyprodukowal przeliczonego pliku.\n"
            f"stdout: {wynik.stdout}\nstderr: {wynik.stderr}"
        )
    return przeliczony


def zbierz_bledy(plik: Path) -> Tuple[List[str], Dict[str, int]]:
    wb = openpyxl.load_workbook(plik, data_only=True)
    bledy: List[str] = []
    policzone: Dict[str, int] = {}
    for nazwa in wb.sheetnames:
        ws = wb[nazwa]
        liczba = 0
        for wiersz in ws.iter_rows():
            for komorka in wiersz:
                if komorka.value is None:
                    continue
                if isinstance(komorka.value, (int, float)):
                    liczba += 1
                    continue
                tekst = str(komorka.value)
                if any(tekst.startswith(b) or tekst == b for b in BLEDY_EXCEL):
                    bledy.append(f"{nazwa}!{komorka.coordinate}: {tekst}")
                elif tekst.startswith("="):
                    bledy.append(
                        f"{nazwa}!{komorka.coordinate}: formula nieprzeliczona ({tekst[:60]})"
                    )
                else:
                    liczba += 1
        policzone[nazwa] = liczba
    return bledy, policzone


def porownaj_z_silnikiem(plik: Path, wejscie: Path) -> List[str]:
    """Kontrola krzyzowa: kluczowe komorki arkusza kontra wynik silnika."""
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from sim_kalkulator.dane import wczytaj_yaml
    from sim_kalkulator.silnik import przelicz as przelicz_silnikiem

    r = przelicz_silnikiem(wczytaj_yaml(wejscie))
    wb = openpyxl.load_workbook(plik, data_only=True)

    def szukaj(arkusz: str, etykieta: str, kolumna: int = 2):
        ws = wb[arkusz]
        for wiersz in ws.iter_rows():
            if wiersz[0].value == etykieta:
                return wiersz[kolumna - 1].value
        return None

    kontrole = [
        ("Alokacja", "Koszty przedsiewziecia", 2, float(r.alokacja.koszty_laczne)),
        ("Alokacja", "GRANT", 3, float(r.granty.spoleczna.kwota)),
        ("Alokacja", "GRANT", 4, float(r.granty.komunalna.kwota)),
        ("Pula_spoleczna", "Wklad wlasny", 2, float(r.finansowanie.spoleczna.wklad_wlasny)),
        ("Pula_komunalna", "Wklad wlasny", 2, float(r.finansowanie.komunalna.wklad_wlasny)),
        ("Pula_spoleczna", "LIMIT WIAZACY", 2, float(r.limity_spoleczna.limit_wiazacy_m2_mies)),
        ("Pula_komunalna", "LIMIT WIAZACY", 2, float(r.limity_komunalna.limit_wiazacy_m2_mies)),
        ("Werdykty", "Wklad wlasny wymagany", 2, float(r.finansowanie.wklad_wlasny_wymagany)),
        ("Werdykty", "LUKA KAPITALOWA", 2, float(r.werdykty.montaz.luka_kwota)),
    ]
    if r.projekcja.spoleczna.minimalny_dscr is not None:
        kontrole.append(
            ("Pula_spoleczna", "Minimalny DSCR", 15, float(r.projekcja.spoleczna.minimalny_dscr))
        )
    if r.edb_kredytu > 0:
        kontrole.append(("Rekompensata", "EDB kredytu", 11, float(r.edb_kredytu)))

    rozjazdy = []
    for arkusz, etykieta, kolumna, oczekiwane in kontrole:
        z_arkusza = szukaj(arkusz, etykieta, kolumna)
        if z_arkusza is None:
            rozjazdy.append(f"{arkusz}: nie znaleziono wiersza '{etykieta}'")
            continue
        if not isinstance(z_arkusza, (int, float)):
            rozjazdy.append(f"{arkusz}!'{etykieta}': wartosc nieliczbowa {z_arkusza!r}")
            continue
        tolerancja = max(abs(oczekiwane) * 1e-6, 0.01)
        if abs(z_arkusza - oczekiwane) > tolerancja:
            rozjazdy.append(
                f"{arkusz}!'{etykieta}': arkusz {z_arkusza:,.2f} vs silnik {oczekiwane:,.2f}"
            )
    return rozjazdy


def main() -> int:
    parser = argparse.ArgumentParser(description="Przeliczenie arkusza i kontrola bledow formul.")
    parser.add_argument("plik", type=Path, help="Skoroszyt XLSX do przeliczenia.")
    parser.add_argument(
        "--porownaj-z-silnikiem", type=Path, default=None,
        help="Plik YAML wejscia — porownuje kluczowe komorki arkusza z wynikiem silnika.",
    )
    args = parser.parse_args()

    if not args.plik.exists():
        print(f"Nie ma pliku {args.plik}", file=sys.stderr)
        return 2

    print(f"Przeliczam {args.plik} przez LibreOffice...")
    przeliczony = przelicz(args.plik)
    bledy, policzone = zbierz_bledy(przeliczony)

    print("\nPrzeliczone komorki w zakladkach:")
    for nazwa, liczba in policzone.items():
        print(f"  {nazwa:20} {liczba:5d}")

    if bledy:
        print(f"\nBLEDY FORMUL: {len(bledy)}", file=sys.stderr)
        for blad in bledy[:60]:
            print(f"  {blad}", file=sys.stderr)
        if len(bledy) > 60:
            print(f"  ... i {len(bledy) - 60} dalszych", file=sys.stderr)
        return 1

    print("\nBledow formul: 0")

    if args.porownaj_z_silnikiem is not None:
        print(f"\nKontrola krzyzowa z silnikiem ({args.porownaj_z_silnikiem})...")
        rozjazdy = porownaj_z_silnikiem(przeliczony, args.porownaj_z_silnikiem)
        if rozjazdy:
            print(f"ROZJAZD ARKUSZ-SILNIK: {len(rozjazdy)}", file=sys.stderr)
            for r in rozjazdy:
                print(f"  {r}", file=sys.stderr)
            return 1
        print("Rozjazdow arkusz-silnik: 0")

    print("\nUWAGA: zielony przebieg dowodzi, ze formuly sie licza, nie ze sa poprawne.")
    print("Sprawdz recznie 2-3 formuly w kazdej zakladce projekcji.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
